# -*- coding: utf-8 -*-
"""
Carga Colombia enero-marzo 2026 al bloque "Consolidacion USD".

POR QUE
-------
La serie "Ingresos totales por mes" muestra s/d en ene, feb y mar de 2026 mientras
tiene cifras en ago-dic 2025 y en abr-jul 2026. La auditoria lo llamo hallazgo I-21 y
lo diagnostico como "el hueco es de cableado, no de datos". No es asi.

El bloque mensual "Consolidacion USD" trae en ene-mar solo a las tres filiales:
Colombia no se cargo ahi hasta abril. El tablero exige los cuatro paises antes de
armar un consolidado, y con razon: sumar lo que hay daria USD 3.499 en enero contra
USD 214.383 en abril -- un consolidado sin Colombia, que es el 98% de la compania.
El s/d es la conducta correcta, y el propio codigo lo explica.

O sea: el hueco es de datos, no de cableado. Esto lo cierra.

CONVERSION
----------
Colombia se origina en COP. Se convierte con la TRM promedio mensual de la hoja TRM,
que es la base que el tablero adopto para el estado de resultados (NIC 21: promedio
del periodo para el resultado, cierre para el balance).

    enero    3.703,00      febrero  3.675,64      marzo  3.717,56

Los ocho codigos son los mismos que ya traen las tres filiales en ese bloque. No se
inventa ninguno: si Colombia no tiene la linea en pesos, esa fila no se crea y se avisa.

La fila "Consolidado" de esos meses NO se crea: el tablero la arma sumando los cuatro
paises cuando no existe (sumaPaisesConsol), y asi queda derivada en vez de guardada,
que es justo lo que evita que se desalinee.

Uso:  python cargar_colombia_ene_mar_usd.py             -> vista previa
      python cargar_colombia_ene_mar_usd.py --escribir  -> aplica, con respaldo
"""
import os
import shutil
import sys
from datetime import datetime

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.join(AQUI, "BD_MAESTRA_COCO.xlsx")
RESPALDOS = os.path.join(AQUI, "respaldos")
HOJA = "BD_Indicadores"

PAIS = "Colombia"
COMPANIA = "Coco Colombia"
SEGMENTO = "Consolidación USD"
MESES = ["2026-01", "2026-02", "2026-03"]
CODIGOS = ["pyg_ingresos_operacionales", "pyg_costo_ventas", "pyg_gasto_administracion",
           "pyg_gasto_ventas", "pyg_gasto_proyectos", "pyg_ingresos_no_operacionales",
           "pyg_gastos_financieros", "pyg_utilidad_neta"]
FUENTE = ("Colombia ene-mar 2026 llevada al bloque USD desde su P&G en COP, con la TRM "
          "promedio mensual de la hoja TRM. Cierra el hueco de la serie consolidada "
          "(hallazgo I-21): el bloque solo traia las filiales.")


def main():
    escribir = "--escribir" in sys.argv

    # data_only=False: el libro tiene 3.882 formulas y guardarlo en modo valores las borra
    wb = openpyxl.load_workbook(BASE, data_only=False)
    ws = wb[HOJA]

    # --- TRM promedio mensual ---
    trm = {}
    wt = wb["TRM"]
    hdr = [str(c.value).strip().lower() if c.value else "" for c in wt[3]]
    try:
        ip, iprom = hdr.index("periodo"), hdr.index("trm_promedio_mercado")
    except ValueError:
        print("No encuentro las columnas periodo / trm_promedio_mercado en la hoja TRM.")
        return 1
    for r in range(4, wt.max_row + 1):
        p, v = wt.cell(r, ip + 1).value, wt.cell(r, iprom + 1).value
        if p and isinstance(v, (int, float)) and v > 0:
            trm[str(p).strip()] = float(v)

    faltan_trm = [m for m in MESES if m not in trm]
    if faltan_trm:
        print("Sin TRM promedio para: %s. No se inventa una tasa; aborta." % ", ".join(faltan_trm))
        return 1

    # --- encabezado ---
    enc = None
    for i, row in enumerate(ws.iter_rows(max_row=6, values_only=True), 1):
        if row and row[0] and str(row[0]).strip() == "periodo":
            enc = (i, [str(x).strip() if x is not None else "" for x in row])
            break
    if not enc:
        print("No encuentro el encabezado.")
        return 1
    fila_enc, cols = enc
    c = {n: cols.index(n) + 1 for n in ("periodo", "pais", "compania", "moneda", "escenario",
                                        "codigo_indicador", "segmento", "valor", "unidad",
                                        "comentario", "area", "area_responsable", "periodicidad")}

    # --- Colombia en COP (segmento vacio) y lo que ya exista en el bloque USD ---
    cop, ya = {}, set()
    for r in range(fila_enc + 1, ws.max_row + 1):
        per = str(ws.cell(r, c["periodo"]).value or "").strip()
        if per not in MESES:
            continue
        pais = str(ws.cell(r, c["pais"]).value or "").strip()
        cod = str(ws.cell(r, c["codigo_indicador"]).value or "").strip()
        seg = str(ws.cell(r, c["segmento"]).value or "").strip()
        esc = str(ws.cell(r, c["escenario"]).value or "").strip()
        if pais == PAIS and seg == "" and esc == "Real":
            cop[(per, cod)] = ws.cell(r, c["valor"]).value
        if pais == PAIS and seg == SEGMENTO:
            ya.add((per, cod))

    nuevas, sin_origen = [], []
    for per in MESES:
        for cod in CODIGOS:
            if (per, cod) in ya:
                continue
            v = cop.get((per, cod))
            if not isinstance(v, (int, float)):
                sin_origen.append((per, cod))
                continue
            nuevas.append((per, cod, v, round(v / trm[per], 2)))

    print("=" * 78)
    print("Colombia ene-mar 2026  ->  bloque '%s'" % SEGMENTO)
    print("=" * 78)
    print("%-9s %-34s %16s %6s %12s" % ("MES", "CODIGO", "COP", "TRM", "USD"))
    print("-" * 78)
    for per, cod, v, usd in nuevas:
        print("%-9s %-34s %16s %6.0f %12s" % (per, cod, "{:,.0f}".format(v).replace(",", "."),
                                              trm[per], "{:,.2f}".format(usd).replace(",", "@")
                                              .replace(".", ",").replace("@", ".")))
    print("-" * 78)
    print("filas a crear: %d" % len(nuevas))
    if ya:
        print("ya existian y no se tocan: %d" % len(ya))
    if sin_origen:
        print("SIN LINEA EN PESOS (no se crean, se avisan): %s"
              % ", ".join("%s/%s" % x for x in sin_origen))

    ing = {p: next((u for (q, cd, _, u) in nuevas if q == p and cd == "pyg_ingresos_operacionales"), None)
           for p in MESES}
    print()
    print("efecto en la serie consolidada (hoy s/d):")
    for p in MESES:
        print("   %s  Colombia USD %s" % (p, "{:,.0f}".format(ing[p]).replace(",", ".") if ing[p] else "—"))

    if not escribir:
        print()
        print("VISTA PREVIA. Nada se escribio. Corre con --escribir para aplicar.")
        return 0

    os.makedirs(RESPALDOS, exist_ok=True)
    sello = datetime.now().strftime("%Y%m%d_%H%M%S")
    resp = os.path.join(RESPALDOS, "BD_MAESTRA_COCO_antes_colombia_ene_mar_%s.xlsx" % sello)
    shutil.copy2(BASE, resp)

    fila = ws.max_row + 1
    for per, cod, _, usd in nuevas:
        ws.cell(fila, c["periodo"]).value = per
        ws.cell(fila, c["pais"]).value = PAIS
        ws.cell(fila, c["compania"]).value = COMPANIA
        ws.cell(fila, c["moneda"]).value = "USD"
        ws.cell(fila, c["escenario"]).value = "Real"
        ws.cell(fila, c["codigo_indicador"]).value = cod
        ws.cell(fila, c["segmento"]).value = SEGMENTO
        ws.cell(fila, c["valor"]).value = usd
        ws.cell(fila, c["unidad"]).value = "USD"
        ws.cell(fila, c["comentario"]).value = FUENTE
        ws.cell(fila, c["area"]).value = "Internacional"
        ws.cell(fila, c["area_responsable"]).value = "Finanzas"
        ws.cell(fila, c["periodicidad"]).value = "Mensual"
        fila += 1

    wb.save(BASE)
    print()
    print("respaldo: %s" % os.path.basename(resp))
    print("creadas %d filas." % len(nuevas))
    print("Corre Revisar_Base.bat: el consolidado de esos meses queda derivado como suma.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
