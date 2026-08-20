# -*- coding: utf-8 -*-
"""
Carga julio 2026 de Cartera & Liquidez, recibido en USD (confirmado) y convertido a
COP con la TRM oficial de julio (3505.2683), igual que el resto de la base.

Quedan SIN cargar, a la espera de confirmar la formula/alcance con quien las envio:
  - indice_morosidad (27.81 en junio -> 77.3 en julio, casi el triple, no coincide
    con vencida/total en ningun mes probado)
  - concentracion_cartera (17.1 -> 45.3, casi el triple, sin explicacion encontrada)

cartera_arr_pct SI se carga, pero con nota: su caida (46.47 -> 20.90) coincide con
usar 'arr_mrr' (el ARR nuevo, MRR x12) como denominador en vez de 'arr' (ventas
totales) -- queda documentado en el comentario, no se asume sin dejarlo dicho.

Uso:  python cargar_cartera_julio.py             -> vista previa
      python cargar_cartera_julio.py --escribir  -> aplica (con respaldo)
"""
import os
import shutil
import sys
from datetime import datetime

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.join(AQUI, "BD_MAESTRA_COCO.xlsx")
TRM_JULIO = 3505.2683
PERIODO = "2026-07"

# (codigo, valor_usd_o_nativo, es_moneda, comentario)
DATOS = [
    ("cartera_total", 469850, True, "Cartera & Liquidez julio 2026 (dato en USD, convertido a COP con TRM julio)."),
    ("cartera_vencida", 195494, True, "Cartera & Liquidez julio 2026 (USD -> COP, TRM julio)."),
    ("cartera_sana", 274355, True, "Cartera & Liquidez julio 2026 (USD -> COP, TRM julio)."),
    ("dso", 58, False, "Cartera & Liquidez julio 2026. Mejoró vs 70,76 días de junio."),
    ("cartera_cobro_legal", 16342, True, "Cartera & Liquidez julio 2026 (USD -> COP, TRM julio)."),
    ("cartera_mayor_60_pct", 7, False, "Cartera & Liquidez julio 2026 (%). Baja fuerte vs 20,3% de junio -- validar con la fuente."),
    ("write_off_ratio", 0.22, False, "Cartera & Liquidez julio 2026."),
    ("cartera_arr_pct", 20.90, False, "Cartera & Liquidez julio 2026 (%). Cae de 46,47 a 20,90 -- consistente con usar "
     "'arr_mrr' (MRR x12) como base en vez de 'arr' (ventas totales); no confirmado con la fuente."),
]


def main():
    escribir = "--escribir" in sys.argv
    wb = openpyxl.load_workbook(BASE)
    ws = wb["BD_Indicadores"]
    H = [str(c.value).strip() if c.value else "" for c in ws[3]]
    ix = {c: H.index(c) + 1 for c in H if c}

    existentes = {}
    for f in range(4, ws.max_row + 1):
        cod = ws.cell(f, ix["codigo_indicador"]).value
        per = ws.cell(f, ix["periodo"]).value
        pais = ws.cell(f, ix["pais"]).value
        if cod and pais == "Colombia":
            existentes[(cod, str(per))] = f

    siguiente = ws.max_row + 1
    reporte = []

    for cod, valor, es_moneda, comentario in DATOS:
        valor_cop = round(valor * TRM_JULIO, 2) if es_moneda else valor
        clave = (cod, PERIODO)
        if clave in existentes:
            f = existentes[clave]
            antes = ws.cell(f, ix["valor"]).value
            reporte.append("  UPDATE %-22s %s -> %s" % (cod, antes, valor_cop))
        else:
            f = siguiente
            siguiente += 1
            ws.cell(f, ix["periodo"]).value = PERIODO
            ws.cell(f, ix["pais"]).value = "Colombia"
            ws.cell(f, ix["compania"]).value = "Coco Colombia"
            ws.cell(f, ix["escenario"]).value = "Real"
            ws.cell(f, ix["periodicidad"]).value = "Mensual"
            reporte.append("  INSERT %-22s -> %s" % (cod, valor_cop))
        ws.cell(f, ix["codigo_indicador"]).value = cod
        ws.cell(f, ix["valor"]).value = valor_cop
        ws.cell(f, ix["moneda"]).value = "COP"
        ws.cell(f, ix["unidad"]).value = "COP" if es_moneda else ("días" if cod == "dso" else "%")
        ws.cell(f, ix["comentario"]).value = comentario

    print("=" * 90)
    print("  CARGA CARTERA & LIQUIDEZ - JULIO 2026 (USD -> COP, TRM %.4f)" % TRM_JULIO)
    print("=" * 90)
    print()
    for l in reporte:
        print(l)
    print("\n  SIN cargar (pendiente de confirmar formula): indice_morosidad, concentracion_cartera\n")

    if not escribir:
        print("  Vista previa. Para aplicar agrega --escribir\n")
        return

    os.makedirs(os.path.join(AQUI, "respaldos"), exist_ok=True)
    sello = datetime.now().strftime("%Y%m%d_%H%M%S")
    copia = os.path.join(AQUI, "respaldos", "BD_MAESTRA_COCO_antes_cartera_julio_" + sello + ".xlsx")
    shutil.copy2(BASE, copia)
    wb.save(BASE)
    print("  Respaldo:", os.path.basename(copia))
    print("  Base actualizada.\n")


if __name__ == "__main__":
    main()
