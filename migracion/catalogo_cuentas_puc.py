# -*- coding: utf-8 -*-
"""
Saca el catalogo de cuentas PUC de gasto de los cuatro paises, con su movimiento
acumulado y la regla de homologacion que ya tiene (o la que le falta).

PARA QUE
--------
La cadena de gastos homologa por SUBCUENTA PUC de nivel 6 (Sueldos, Asesoria juridica,
Viaticos...). Las filiales estan homologadas; Colombia entro como totales --PL-O-200
"Total administracion", PL-S-200 "Total ventas", PL-O-210 "Total operativo"-- no porque
le falte detalle, sino porque asi se cargo. Tiene 99 subcuentas de gasto en su auxiliar.

De esas 99, solo 30 tienen regla en EEFF_Homologacion. Las 90 reglas existentes se
escribieron mirando a las filiales, que usan muchas menos cuentas.

Este script arma la lista de lo que falta, ordenada por plata, para que contabilidad
decida a que codigo comun va cada una. Con eso Colombia se homologa igual que las
filiales y el grafico de gastos deja de ser "consolidado parcial".

SALIDA
------
migracion/Catalogo_cuentas_PUC.xlsx con dos hojas:

  Cuentas    una fila por subcuenta y pais, con codigo PUC, jerarquia, movimiento
             acumulado en moneda de origen, y el codigo comun si ya lo tiene
  Faltantes  solo las que NO tienen regla, ordenadas por movimiento descendente,
             con la columna codigo_comun_propuesto EN BLANCO para diligenciar

Uso:  python catalogo_cuentas_puc.py
"""
import io
import os
import sys
import zipfile
from collections import OrderedDict

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

AQUI = os.path.dirname(os.path.abspath(__file__))
TRAZA = os.path.join(AQUI, "BD_TRAZABILIDAD_COCO.xlsx")
SALIDA = os.path.join(AQUI, "Catalogo_cuentas_PUC.xlsx")

# Los auxiliares tal como los entrega el ERP. Se aceptan las dos convenciones de
# nombre que conviven hoy ("Julio 2026 USA.xlsx" y "Julio Peru 2026.xlsx").
CARPETAS = [
    os.path.join(os.path.expanduser("~"), "Desktop", "Temporales"),
    AQUI,
]
MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio"]
PAISES = OrderedDict([
    ("Colombia", ["{m} Colombia 2026.xlsx"]),
    ("EE.UU.",   ["{m} 2026 USA.xlsx"]),
    # OJO con Peru: el orden importa. "{m} Peru 2026.xlsx" es el zip corregido del
    # 27-ago; "{m} 2026 PERU.xlsx" es el anterior, cuyo julio era junio reexportado.
    ("Perú",     ["{m} Peru 2026.xlsx", "{m} 2026 PERU.xlsx"]),
    ("Costa Rica", ["Costa Rica {m} 2026 .xlsx", "{m} 2026 CR.xlsx"]),
])
MONEDA = {"Colombia": "COP", "EE.UU.": "USD", "Perú": "PEN", "Costa Rica": "USD"}

NIVEL = {1: "Clase", 2: "Grupo", 4: "Cuenta", 6: "Subcuenta"}


def buscar(nombre):
    """Devuelve el auxiliar como bytes. Lo busca suelto y TAMBIEN dentro de los .zip
    que entrega el ERP, que es como llegan de verdad ("EF Colombia.zip",
    "Fwd_ Estados financieros Mensuales Filiales.zip", "EF Peru.zip")."""
    for c in CARPETAS:
        if not os.path.isdir(c):
            continue
        for raiz, _, files in os.walk(c):
            if nombre in files:
                return io.BytesIO(open(os.path.join(raiz, nombre), "rb").read())
            for f in files:
                if not f.lower().endswith(".zip"):
                    continue
                try:
                    z = zipfile.ZipFile(os.path.join(raiz, f))
                except Exception:
                    continue
                for i in z.namelist():
                    if os.path.basename(i) == nombre:
                        return io.BytesIO(z.read(i))
    return None


def leer_auxiliar(ruta):
    """Devuelve {codigo_puc: (nombre, movimiento_neto)} para las cuentas de gasto."""
    try:
        wb = openpyxl.load_workbook(ruta, data_only=True)
    except Exception:
        return None            # texto tabulado con extension .xlsx: se avisa, no se adivina
    out = {}
    for fila in wb.active.iter_rows(values_only=True):
        if not fila:
            continue
        etq = None
        for x in fila[:3]:
            if isinstance(x, str) and " - " in x:
                etq = x.strip()
                break
        if not etq:
            continue
        cod = etq.split(" - ")[0].strip()
        if not cod.isdigit() or cod[0] != "5":
            continue
        nums = [v for v in fila if isinstance(v, (int, float))]
        if len(nums) < 4:
            continue
        deb, cre = nums[-3], nums[-2]
        nom = etq.split(" - ", 1)[1].strip()
        a, b = out.get(cod, (nom, 0.0))
        out[cod] = (a, b + (deb - cre))
    return out


def reglas_homologacion():
    wb = openpyxl.load_workbook(TRAZA, data_only=True)
    filas = list(wb["EEFF_Homologacion"].iter_rows(values_only=True))
    cab = [str(v or "").strip() for v in filas[2]]
    ix = {n: cab.index(n) for n in ("Código fuente", "Código común", "Rubro común", "Subrubro")}
    reglas = {}
    for f in filas[3:]:
        if not f or not f[ix["Código fuente"]]:
            continue
        reglas[str(f[ix["Código fuente"]]).strip()] = (
            str(f[ix["Código común"]] or ""), str(f[ix["Rubro común"]] or ""),
            str(f[ix["Subrubro"]] or ""))
    return reglas


def main():
    reglas = reglas_homologacion()
    print("reglas de homologacion existentes: %d" % len(reglas))

    datos, ilegibles = {}, []
    for pais, patrones in PAISES.items():
        for mes in MESES:
            ruta = None
            for pat in patrones:
                ruta = buscar(pat.format(m=mes))
                if ruta:
                    break
            if not ruta:
                continue
            aux = leer_auxiliar(ruta)
            if aux is None:
                ilegibles.append((pais, mes, pat.format(m=mes)))
                continue
            for cod, (nom, mov) in aux.items():
                k = (pais, cod)
                a, b = datos.get(k, (nom, 0.0))
                datos[k] = (a or nom, b + mov)

    print("archivos ilegibles (texto tabulado con extension .xlsx): %d" % len(ilegibles))
    for p, m, n in ilegibles:
        print("   %-12s %-9s %s" % (p, m, n))

    filas = []
    for (pais, cod), (nom, mov) in datos.items():
        r = reglas.get(cod)
        filas.append(OrderedDict([
            ("pais", pais),
            ("moneda", MONEDA[pais]),
            ("cuenta_puc", cod),
            ("nivel", NIVEL.get(len(cod), "n%d" % len(cod))),
            ("nombre_cuenta", nom),
            ("clase", cod[:1]),
            ("grupo", cod[:2]),
            ("cuenta", cod[:4]),
            ("movimiento_ene_jul", round(mov, 2)),
            ("codigo_comun", r[0] if r else ""),
            ("rubro_comun", r[1] if r else ""),
            ("subrubro", r[2] if r else ""),
            ("tiene_regla", "SI" if r else "NO"),
        ]))
    # solo el nivel de detalle que homologa la cadena
    detalle = [f for f in filas if len(f["cuenta_puc"]) == 6]
    detalle.sort(key=lambda f: (f["pais"], f["cuenta_puc"]))
    faltan = sorted([f for f in detalle if f["tiene_regla"] == "NO"],
                    key=lambda f: -abs(f["movimiento_ene_jul"]))

    wb = openpyxl.Workbook()
    hdrf = Font(bold=True, color="FFFFFF")
    hdrb = PatternFill("solid", fgColor="1D283A")
    azul = PatternFill("solid", fgColor="DCE9F7")

    def hoja(ws, cols, rows, editable=None):
        ws.append(cols)
        for c in ws[1]:
            c.font = hdrf
            c.fill = hdrb
            c.alignment = Alignment(horizontal="center")
        for r in rows:
            ws.append([r.get(c) for c in cols])
        for i, c in enumerate(cols, 1):
            ancho = max(11, min(38, max([len(str(c))] + [len(str(r.get(c) or ""))
                                                         for r in rows[:400]]) + 2))
            ws.column_dimensions[ws.cell(1, i).column_letter].width = ancho
            if editable and c in editable:
                for j in range(2, len(rows) + 2):
                    ws.cell(j, i).fill = azul
        ws.freeze_panes = "A2"

    if not detalle:
        print("No encontre ningun auxiliar. Deja los .zip del ERP en Escritorio\Temporales")
        print("o en migracion/, y vuelve a correr.")
        return 1
    ws1 = wb.active
    ws1.title = "Cuentas"
    hoja(ws1, list(detalle[0].keys()), detalle)

    ws2 = wb.create_sheet("Faltantes")
    cols2 = ["pais", "moneda", "cuenta_puc", "nombre_cuenta", "cuenta", "movimiento_ene_jul",
             "codigo_comun_propuesto", "notas"]
    for f in faltan:
        f["codigo_comun_propuesto"] = ""
        f["notas"] = ""
    hoja(ws2, cols2, faltan, editable=("codigo_comun_propuesto", "notas"))

    wb.save(SALIDA)

    print()
    print("=" * 72)
    print("subcuentas de gasto (nivel 6) encontradas: %d" % len(detalle))
    print("  con regla : %d" % sum(1 for f in detalle if f["tiene_regla"] == "SI"))
    print("  SIN regla : %d" % len(faltan))
    print()
    print("por pais:")
    for pais in PAISES:
        d = [f for f in detalle if f["pais"] == pais]
        n = sum(1 for f in d if f["tiene_regla"] == "NO")
        print("   %-12s %3d subcuentas · %3d sin regla" % (pais, len(d), n))
    print()
    print("archivo: %s" % os.path.basename(SALIDA))
    return 0


if __name__ == "__main__":
    sys.exit(main())
