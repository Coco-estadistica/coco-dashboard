# -*- coding: utf-8 -*-
"""
Carga el movimiento mensual de gastos de Colombia a BD_GASTOS_HOMOLOGADOS.

POR QUE
-------
La cadena de gastos cubre a las tres filiales de enero a julio, pero de Colombia solo
tiene doce filas de julio con el ACUMULADO cargado y el movimiento mensual en CERO. Por
eso "Gastos Consolidado" nunca sumo a Colombia, que es el 98% de la compania, y el HTML
tuvo que parchearlo inyectando el detalle real (README 4.2, monthlyGastosBridgeData).

Esto carga el movimiento mes a mes, que es lo que faltaba.

EL MAPEO NO ES INVENTADO
------------------------
Cada codigo comun corresponde a una linea del P&G mensual de Colombia, y la equivalencia
la confirma la propia base: el acumulado que YA estaba cargado en la cadena coincide al
peso con la suma enero-julio de esa linea.

    PL-C-010  Costo de ventas          <- pyg_costo_ventas          2.442.623.347
    PL-O-200  Gastos de administracion <- pyg_gasto_administracion    797.478.918
    PL-S-200  Gastos de ventas         <- pyg_gasto_ventas            959.736.563
    PL-O-210  Gastos operacionales     <- pyg_gasto_proyectos         676.358.287
    PL-F-010  Gastos financieros       <- pyg_gastos_financieros       40.108.353

LO QUE ESTO NO RESUELVE
----------------------
Colombia aporta TOTALES por linea contable; las filiales aportan DETALLE (personal,
honorarios, viajes, seguros...). Son dos taxonomias distintas y esa diferencia no se
salva cargando datos:

  - el area de Colombia vive en gasto_area, con cinco rubros de gestion
    (Financieros, TI, Proveedores, Comercial, Administracion), 31 meses desde 2024-01
  - el area de las filiales sale del mapeo contable homologado, 7 meses

Con esto el total por pais queda bien y el consolidado suma los cuatro. El desglose por
area de Colombia sigue viniendo de gasto_area, no de esta cadena. Unificar las dos
taxonomias es una decision de negocio, no un cargue.

Uso:  python cargar_gastos_colombia_cadena.py             -> vista previa
      python cargar_gastos_colombia_cadena.py --escribir  -> aplica, con respaldo
"""
import os
import shutil
import sys
from datetime import datetime

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.join(AQUI, "BD_MAESTRA_COCO.xlsx")
RESPALDOS = os.path.join(AQUI, "respaldos")
HOJA = "BD_GASTOS_HOMOLOGADOS"

PAIS, ENTIDAD, MONEDA = "Colombia", "COCO", "COP"
MESES = ["2026-0%d" % i for i in range(1, 8)]

# codigo_comun -> (rubro_comun, categoria, codigo del P&G mensual de Colombia)
MAPEO = [
    ("PL-C-010", "Costo de ventas", "Costo de ventas", "pyg_costo_ventas"),
    ("PL-O-200", "Gastos de administración - COCO", "Gastos de administración", "pyg_gasto_administracion"),
    ("PL-S-200", "Gastos de ventas - COCO", "Gastos de ventas", "pyg_gasto_ventas"),
    ("PL-O-210", "Gastos operacionales - COCO", "Otros gastos operacionales", "pyg_gasto_proyectos"),
    ("PL-F-010", "Gastos financieros", "Gastos financieros", "pyg_gastos_financieros"),
]
FUENTE = "P&G mensual oficial de Colombia (BD_Indicadores, segmento vacío)"
NOTA = ("Movimiento mensual cargado desde el P&G oficial de Colombia. La equivalencia "
        "codigo comun <-> linea del P&G esta verificada contra el acumulado que ya "
        "estaba en esta hoja: coincide al peso en las cinco.")
TOL = 2.0


def main():
    escribir = "--escribir" in sys.argv
    wb = openpyxl.load_workbook(BASE, data_only=False)   # conserva las 3.882 formulas

    # --- P&G mensual de Colombia ---
    ws = wb["BD_Indicadores"]
    enc = None
    for i, row in enumerate(ws.iter_rows(max_row=6, values_only=True), 1):
        if row and row[0] and str(row[0]).strip() == "periodo":
            enc = (i, [str(x).strip() if x is not None else "" for x in row])
            break
    fi, cols = enc
    ci = {n: cols.index(n) for n in ("periodo", "pais", "escenario", "codigo_indicador",
                                     "segmento", "valor")}
    pyg = {}
    for row in ws.iter_rows(min_row=fi + 1, values_only=True):
        if not row or not row[ci["periodo"]]:
            continue
        if (str(row[ci["pais"]] or "").strip() != PAIS
                or str(row[ci["escenario"]] or "").strip() != "Real"
                or str(row[ci["segmento"]] or "").strip() != ""):
            continue
        pyg[(str(row[ci["periodo"]]).strip(), str(row[ci["codigo_indicador"]]).strip())] = row[ci["valor"]]

    # --- hoja de gastos ---
    wg = wb[HOJA]
    hg = [str(c.value).strip() if c.value else "" for c in wg[1]]
    gi = {n: hg.index(n) for n in hg if n}
    existentes = {}
    for r in range(2, wg.max_row + 1):
        if str(wg.cell(r, gi["pais"] + 1).value or "").strip() != PAIS:
            continue
        existentes[(str(wg.cell(r, gi["periodo"] + 1).value or "").strip(),
                    str(wg.cell(r, gi["codigo_comun"] + 1).value or "").strip())] = r

    print("=" * 96)
    print("Colombia -> BD_GASTOS_HOMOLOGADOS   (movimiento mensual)")
    print("=" * 96)

    # control: el acumulado ya cargado debe coincidir con la suma ene-jul
    print("%-10s %-30s %18s %18s %s" % ("CODIGO", "LINEA DEL P&G", "YTD YA CARGADO", "SUMA ENE-JUL", "CONTROL"))
    print("-" * 96)
    f = lambda v: "{:,.0f}".format(v or 0).replace(",", ".")
    aborta = False
    for cod, rubro, cat, code in MAPEO:
        suma = sum(pyg.get((m, code)) or 0 for m in MESES)
        r = existentes.get(("2026-07", cod))
        ytd = wg.cell(r, gi["saldo_presentacion_ytd"] + 1).value if r else None
        ok = ytd is not None and abs((ytd or 0) - suma) <= TOL
        print("%-10s %-30s %18s %18s %s" % (cod, code, f(ytd), f(suma), "OK" if ok else "NO CUADRA"))
        if not ok:
            aborta = True
    if aborta:
        print("\nABORTA: el mapeo no reproduce el acumulado ya cargado. No se escribe nada.")
        return 1

    nuevas, actualizadas = [], []
    for cod, rubro, cat, code in MAPEO:
        acum = 0.0
        for m in MESES:
            v = pyg.get((m, code))
            if not isinstance(v, (int, float)):
                continue
            acum += v
            key = (m, cod)
            if key in existentes:
                actualizadas.append((existentes[key], m, cod, v, acum))
            else:
                nuevas.append((m, cod, rubro, cat, v, acum))

    print()
    print("filas nuevas: %d · filas actualizadas (movimiento estaba en cero): %d"
          % (len(nuevas), len(actualizadas)))
    print()
    print("%-10s %-10s %18s" % ("MES", "CODIGO", "MOVIMIENTO COP"))
    for m, cod, _, _, v, _ in nuevas[:6]:
        print("%-10s %-10s %18s" % (m, cod, f(v)))
    if len(nuevas) > 6:
        print("   … y %d filas mas" % (len(nuevas) - 6))

    if not escribir:
        print()
        print("VISTA PREVIA. Nada se escribio. Corre con --escribir para aplicar.")
        return 0

    os.makedirs(RESPALDOS, exist_ok=True)
    sello = datetime.now().strftime("%Y%m%d_%H%M%S")
    resp = os.path.join(RESPALDOS, "BD_MAESTRA_COCO_antes_gastos_colombia_cadena_%s.xlsx" % sello)
    shutil.copy2(BASE, resp)

    for r, m, cod, v, acum in actualizadas:
        wg.cell(r, gi["movimiento_mes"] + 1).value = v
        wg.cell(r, gi["saldo_presentacion_ytd"] + 1).value = acum
        wg.cell(r, gi["fundamento"] + 1).value = NOTA

    fila = wg.max_row + 1
    for m, cod, rubro, cat, v, acum in nuevas:
        wg.cell(fila, gi["periodo"] + 1).value = m
        wg.cell(fila, gi["pais"] + 1).value = PAIS
        wg.cell(fila, gi["entidad"] + 1).value = ENTIDAD
        wg.cell(fila, gi["moneda"] + 1).value = MONEDA
        wg.cell(fila, gi["codigo_comun"] + 1).value = cod
        wg.cell(fila, gi["categoria"] + 1).value = cat
        wg.cell(fila, gi["rubro_comun"] + 1).value = rubro
        wg.cell(fila, gi["movimiento_mes"] + 1).value = v
        wg.cell(fila, gi["saldo_presentacion_ytd"] + 1).value = acum
        wg.cell(fila, gi["fuente_archivo"] + 1).value = FUENTE
        wg.cell(fila, gi["confianza"] + 1).value = "Alta"
        wg.cell(fila, gi["fundamento"] + 1).value = NOTA
        fila += 1

    wb.save(BASE)
    print()
    print("respaldo: %s" % os.path.basename(resp))
    print("escritas %d filas nuevas y %d actualizadas." % (len(nuevas), len(actualizadas)))
    print("SIGUIENTE: regenerar_gastos.py --probar, y Revisar_Base.bat.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
