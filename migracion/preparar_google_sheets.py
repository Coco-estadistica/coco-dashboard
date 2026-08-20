# -*- coding: utf-8 -*-
"""
Prepara el libro que se sube a Google Sheets.

Toma la base vigente (BD_COCO_2026.xlsx) y le agrega las tres hojas de control que
la capa de escritura (Apps Script) necesita y que la base no tiene:

  CATALOGOS          qué códigos puede cargar cada área (sale de las 8 plantillas)
  LOG                bitácora de auditoría (vacía; la llena el Apps Script)
  Pipeline_Comercial oportunidades comerciales (vacía; la llena Comercial)

No modifica la base original: escribe un archivo aparte, BD_MAESTRA_COCO.xlsx.

Uso:  python preparar_google_sheets.py
"""
import os
import shutil

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
ORIGEN = os.path.join(AQUI, "BD_COCO_2026.xlsx")
SALIDA = os.path.join(AQUI, "BD_MAESTRA_COCO.xlsx")
PLANTILLAS = os.path.join(AQUI, "..", "plantillas")

AREAS_PLANTILLA = {
    "Finanzas":         "Plantilla_Carga_Finanzas.xlsx",
    "Ingresos_MRR":     "Plantilla_Carga_Ingresos_MRR.xlsx",
    "Comercial_Ventas": "Plantilla_Carga_Comercial_Ventas.xlsx",
    "Marketing":        "Plantilla_Carga_Marketing.xlsx",
    "Customer_Success": "Plantilla_Carga_Customer_Success.xlsx",
    "Cartera_Liquidez": "Plantilla_Carga_Cartera_Liquidez.xlsx",
    "Costos_Gastos":    "Plantilla_Carga_Costos_Gastos.xlsx",
    "Internacional":    "Plantilla_Carga_Internacional.xlsx",
}

PIPELINE_HEADER = ["cliente", "comercial", "tipo", "prob", "mes",
                   "mrr_cop", "impl_cop", "citas", "valor_cop", "ponderado_cop"]
LOG_HEADER = ["timestamp", "usuario", "accion", "area", "hoja", "filas_afectadas", "detalle"]

PAISES = ["Colombia", "EE.UU.", "Perú", "Costa Rica", "Consolidado", "Otros"]
ESCENARIOS = ["Real", "Presupuesto", "Modelo"]
MONEDAS = ["COP", "USD"]
COMPANIAS = ["Coco Colombia", "Coco LLC", "Coco Perú", "Coco Costa Rica", "Coco Consolidado"]


def codigos_de(plantilla):
    """Lee la hoja 'Listas' de una plantilla: columna A = códigos permitidos del área."""
    wb = openpyxl.load_workbook(plantilla, data_only=True)
    if "Listas" not in wb.sheetnames:
        return []
    return [str(r[0]).strip() for r in wb["Listas"].iter_rows(min_col=1, max_col=1, values_only=True)
            if r[0] is not None and str(r[0]).strip()]


def main():
    if not os.path.exists(ORIGEN):
        raise SystemExit(f"No encuentro la base: {ORIGEN}")

    shutil.copy2(ORIGEN, SALIDA)          # se trabaja sobre una copia
    wb = openpyxl.load_workbook(SALIDA)
    creadas = []

    # --- CATALOGOS: área -> códigos permitidos + listas de validación ---
    if "CATALOGOS" in wb.sheetnames:
        del wb["CATALOGOS"]
    ws = wb.create_sheet("CATALOGOS")
    ws.append(["area", "codigo_indicador"])
    total = 0
    for area, nombre in AREAS_PLANTILLA.items():
        ruta = os.path.join(PLANTILLAS, nombre)
        if not os.path.exists(ruta):
            print(f"   aviso: falta la plantilla {nombre} — el área {area} queda sin códigos")
            continue
        for cod in codigos_de(ruta):
            ws.append([area, cod])
            total += 1
    for col, valores, titulo in (("D", PAISES, "paises"), ("E", ESCENARIOS, "escenarios"),
                                 ("F", MONEDAS, "monedas"), ("G", COMPANIAS, "companias")):
        ws[f"{col}1"] = titulo
        for i, v in enumerate(valores, start=2):
            ws[f"{col}{i}"] = v
    creadas.append(f"CATALOGOS ({total} pares área→código)")

    # --- LOG: bitácora de auditoría ---
    if "LOG" not in wb.sheetnames:
        wb.create_sheet("LOG").append(LOG_HEADER)
        creadas.append("LOG (vacía)")

    # --- Pipeline_Comercial ---
    if "Pipeline_Comercial" not in wb.sheetnames:
        wb.create_sheet("Pipeline_Comercial").append(PIPELINE_HEADER)
        creadas.append("Pipeline_Comercial (vacía)")

    wb.save(SALIDA)

    print("Libro listo para Google Sheets:")
    print(f"   {os.path.abspath(SALIDA)}\n")
    print("Hojas agregadas:")
    for c in creadas:
        print(f"   + {c}")
    wb2 = openpyxl.load_workbook(SALIDA, data_only=True)
    print(f"\nTotal de hojas: {len(wb2.sheetnames)}")
    print("   " + " · ".join(wb2.sheetnames))


if __name__ == "__main__":
    main()
