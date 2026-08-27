# -*- coding: utf-8 -*-
"""
Inventario COMPLETO de variables del tablero: contables y no contables.

QUE INCLUYE
-----------
Todo lo que el tablero conoce, venga de donde venga:

  1. Los 107 indicadores del Diccionario, con su ficha
  2. Los que tienen datos cargados en BD_Indicadores aunque no tengan ficha
  3. Los DERIVADOS, que no se cargan: los calcula el HTML (margen EBITDA, NPS...)
  4. Las 145 subcuentas PUC de gasto de los cuatro paises, que son la capa contable
     de donde salen los indicadores de costos

Para cada una dice de donde viene, quien la carga, que cobertura tiene hoy y hasta
que mes llega. Es el mapa que faltaba para saber que se pide cada cierre y a quien.

SALIDA
------
migracion/Catalogo_variables.xlsx

  Variables    una fila por indicador, con ficha + cobertura real
  Cuentas_PUC  la capa contable: subcuenta, pais, movimiento y su codigo comun
  Resumen      cuantas hay de cada tipo y cuantas estan al dia

Uso:  python catalogo_variables.py
"""
import io
import os
import re
import sys
import zipfile
from collections import OrderedDict, Counter

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

AQUI = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.join(AQUI, "BD_MAESTRA_COCO.xlsx")
HTML = os.path.join(os.path.dirname(AQUI), "Dashboard_COCO.html")
PUC = os.path.join(AQUI, "Catalogo_cuentas_PUC.xlsx")
SALIDA = os.path.join(AQUI, "Catalogo_variables.xlsx")

CORTE = "2026-07"


def norm(s):
    return str(s or "").strip()


def leer_diccionario(wb):
    filas = list(wb["Diccionario"].iter_rows(values_only=True))
    cab = [norm(v) for v in filas[2]]
    ix = {c: i for i, c in enumerate(cab) if c}
    out = {}
    for f in filas[3:]:
        if not f:
            continue
        cod = norm(f[ix["Código indicador"]]) if "Código indicador" in ix else ""
        if not cod:
            continue
        g = lambda c: norm(f[ix[c]]) if c in ix and ix[c] < len(f) else ""
        out[cod] = OrderedDict([
            ("nombre", g("Nombre visible")),
            ("categoria", g("Categoría")),
            ("tipo", g("Tipo")),
            ("formula", g("Fórmula / regla de cálculo")),
            ("unidad", g("Unidad")),
            ("responsable_carga", g("Responsable carga")),
            ("fuente_sugerida", g("Fuente sugerida")),
            ("frecuencia", g("Frecuencia")),
            ("desglose", g("Desglose (segmento / detalle)")),
            ("origen_ficha", g("Origen")),
        ])
    return out


def leer_derivados():
    """Los que NO se cargan: los calcula el tablero. Se leen del propio HTML para que
    la lista no se desactualice cuando alguien agregue o quite uno."""
    try:
        src = io.open(HTML, encoding="utf-8").read()
    except Exception:
        return {}
    i = src.find("const DERIVE=")
    if i < 0:
        return {}
    bloque = src[i:src.find("\n};", i)]
    return {m.group(1): True for m in re.finditer(r"^\s{2}([a-z_0-9]+):", bloque, re.M)}


def main():
    wb = openpyxl.load_workbook(BASE, data_only=True)
    dic = leer_diccionario(wb)
    derivados = leer_derivados()

    filas = [f for f in wb["BD_Indicadores"].iter_rows(values_only=True)][3:]
    cab = [norm(v) for v in list(wb["BD_Indicadores"].iter_rows(values_only=True))[2]]
    ix = {c: i for i, c in enumerate(cab) if c}

    cob = {}
    for f in filas:
        if not f or not f[0]:
            continue
        cod = norm(f[ix["codigo_indicador"]])
        d = cob.setdefault(cod, {"periodos": set(), "paises": set(), "monedas": set(),
                                 "segmentos": set(), "detalles": set(), "escenarios": set(),
                                 "area": "", "resp": "", "periodicidad": "", "n": 0})
        d["n"] += 1
        d["periodos"].add(norm(f[ix["periodo"]]))
        d["paises"].add(norm(f[ix["pais"]]))
        d["monedas"].add(norm(f[ix["moneda"]]))
        d["escenarios"].add(norm(f[ix["escenario"]]))
        if norm(f[ix["segmento"]]):
            d["segmentos"].add(norm(f[ix["segmento"]]))
        if norm(f[ix["detalle"]]):
            d["detalles"].add(norm(f[ix["detalle"]]))
        for k, c in (("area", "area"), ("resp", "area_responsable"), ("periodicidad", "periodicidad")):
            if c in ix and ix[c] < len(f) and norm(f[ix[c]]):
                d[k] = norm(f[ix[c]])

    todos = sorted(set(dic) | set(cob) | set(derivados))
    out = []
    for cod in todos:
        d = dic.get(cod, {})
        c = cob.get(cod)
        pers = sorted(c["periodos"]) if c else []
        ultimo = pers[-1] if pers else ""
        if cod in derivados:
            procede = "Derivado — lo calcula el tablero"
        elif c and any(s == "Consolidación USD" or s.startswith("Acumulado") for s in c["segmentos"]):
            procede = "Contable — auxiliar PUC homologado"
        elif c:
            procede = "Cargue manual / plantilla de área"
        else:
            procede = "Sin datos — solo ficha"
        out.append(OrderedDict([
            ("codigo", cod),
            ("nombre", d.get("nombre", "") or cod),
            ("categoria", d.get("categoria", "") or (c["area"] if c else "")),
            ("procedencia", procede),
            ("unidad", d.get("unidad", "") or (", ".join(sorted(c["monedas"])) if c else "")),
            ("responsable_carga", d.get("responsable_carga", "") or (c["resp"] if c else "")),
            ("fuente_sugerida", d.get("fuente_sugerida", "")),
            ("frecuencia", d.get("frecuencia", "") or (c["periodicidad"] if c else "")),
            ("tiene_ficha", "SI" if cod in dic else "NO"),
            ("tiene_datos", "SI" if c else "NO"),
            ("filas", c["n"] if c else 0),
            ("meses", len(pers)),
            ("primer_mes", pers[0] if pers else ""),
            ("ultimo_mes", ultimo),
            ("al_dia", ("SI" if ultimo >= CORTE else "NO") if ultimo else ""),
            ("paises", ", ".join(sorted(c["paises"])) if c else ""),
            ("escenarios", ", ".join(sorted(c["escenarios"])) if c else ""),
            ("segmentos", ", ".join(sorted(c["segmentos"])) if c else ""),
            ("detalles", ", ".join(sorted(c["detalles"]))[:120] if c else ""),
            ("formula", d.get("formula", "")),
        ]))

    puc = []
    if os.path.exists(PUC):
        wp = openpyxl.load_workbook(PUC, data_only=True)
        fp = list(wp["Cuentas"].iter_rows(values_only=True))
        cp = [norm(v) for v in fp[0]]
        puc = [OrderedDict(zip(cp, f)) for f in fp[1:] if f and f[0]]

    wbo = openpyxl.Workbook()
    hf, hb = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1D283A")

    def hoja(ws, rows, cols=None):
        cols = cols or list(rows[0].keys())
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = hf, hb
            c.alignment = Alignment(horizontal="center")
        for r in rows:
            ws.append([r.get(c) for c in cols])
        for i, c in enumerate(cols, 1):
            w = max(10, min(40, max([len(str(c))] + [len(str(r.get(c) or "")) for r in rows[:300]]) + 2))
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    hoja(wbo.active, out)
    wbo.active.title = "Variables"
    if puc:
        hoja(wbo.create_sheet("Cuentas_PUC"), puc)

    c = Counter(r["procedencia"] for r in out)
    res = [OrderedDict([("concepto", k), ("cantidad", v)]) for k, v in c.most_common()]
    res += [
        OrderedDict([("concepto", "— TOTAL variables"), ("cantidad", len(out))]),
        OrderedDict([("concepto", "con ficha en el Diccionario"), ("cantidad", sum(1 for r in out if r["tiene_ficha"] == "SI"))]),
        OrderedDict([("concepto", "SIN ficha (tienen datos, no estan documentadas)"), ("cantidad", sum(1 for r in out if r["tiene_ficha"] == "NO" and r["tiene_datos"] == "SI"))]),
        OrderedDict([("concepto", "con datos"), ("cantidad", sum(1 for r in out if r["tiene_datos"] == "SI"))]),
        OrderedDict([("concepto", "al dia a " + CORTE), ("cantidad", sum(1 for r in out if r["al_dia"] == "SI"))]),
        OrderedDict([("concepto", "con datos pero rezagadas"), ("cantidad", sum(1 for r in out if r["al_dia"] == "NO"))]),
        OrderedDict([("concepto", "subcuentas PUC de gasto"), ("cantidad", len(puc))]),
    ]
    hoja(wbo.create_sheet("Resumen"), res)
    wbo.save(SALIDA)

    print("=" * 66)
    print("INVENTARIO DE VARIABLES")
    print("=" * 66)
    for r in res:
        print("  %-48s %5s" % (r["concepto"], r["cantidad"]))
    print()
    print("archivo: %s" % os.path.basename(SALIDA))
    return 0


if __name__ == "__main__":
    sys.exit(main())
