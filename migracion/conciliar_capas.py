# -*- coding: utf-8 -*-
"""
CONCILIACIÓN DE CAPAS — control de integridad de la base

Por qué existe
--------------
La misma cifra vive en varias hojas:

  BD_Indicadores    <- la que LEE el dashboard   (fuente de verdad para el tablero)
  BD_PYG_OFICIAL    <- soporte financiero / trazabilidad
  Subsidiarias_PyG  <- resumen por subsidiaria

Si se actualiza una y no las otras, el tablero muestra cifras viejas sin avisar.
Eso ya pasó con Perú julio: estaba en el consolidado estructurado pero en cero en la
hoja que el dashboard lee. Este script detecta ese desalineamiento ANTES de publicar.

Nota (13-ago-2026): tras esa fecha, B y C muestran decenas de diferencias EN
EE.UU./Costa Rica/Consolidado ene-jul a propósito. Se reemplazó el P&G de esas tres
filiales por cifras homologadas desde sus auxiliares contables reales (antes venían
mezcladas: del modelo financiero en abr-jul, sin dato en ene-mar). Subsidiarias_PyG y
BD_PYG_OFICIAL son archivo histórico y no se reescribieron con las cifras nuevas — así
que quedan desalineadas por diseño, no por descuido. El check A (Países vs Consolidado)
es el que importa aquí y debe seguir en OK.

Qué revisa
----------
  A. Países vs Consolidado: los 4 países deben sumar el bloque Consolidado, mes a mes.
  B. BD_Indicadores vs Subsidiarias_PyG: misma cifra por periodo/pais/codigo.
  C. BD_Indicadores vs BD_PYG_OFICIAL: ídem.
  D. Ceros sospechosos: un país con TODO en cero en un mes suele ser dato no cargado,
     no una operación real en cero (regla del proyecto: faltante != cero).

Uso:  python conciliar_capas.py
Salida: OK / diferencias, con el detalle para poder corregir.
No modifica nada.
"""
import hashlib
import json
import os
import sys

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.join(AQUI, "BD_MAESTRA_COCO.xlsx")
# Tras separar la base, las capas de soporte viven en el libro de trazabilidad.
# Se siguen conciliando: si dejaran de compararse, el control perdería sentido.
TRAZA = os.path.join(AQUI, "BD_TRAZABILIDAD_COCO.xlsx")
TOL = 1.0          # tolerancia en USD: redondeos de la fuente
SEG = "Consolidación USD"
PAISES = ["EE.UU.", "Perú", "Costa Rica", "Colombia"]
RUBROS = ["pyg_ingresos_operacionales", "pyg_costo_ventas", "pyg_gasto_administracion",
          "pyg_gasto_ventas", "pyg_gasto_proyectos", "pyg_gastos_financieros",
          "pyg_utilidad_neta"]


def hoja_dict(ws, campos):
    """Lee una hoja con encabezado detectado y devuelve lista de dicts."""
    hdr = None
    for r in range(1, 8):
        fila = [str(c.value).strip().lower() if c.value else "" for c in ws[r]]
        if "periodo" in fila and "codigo_indicador" in fila:
            hdr = r
            break
    if hdr is None:
        return []
    H = [(str(c.value).strip() if c.value else "") for c in ws[hdr]]
    idx = {c: H.index(c) for c in campos if c in H}
    out = []
    for fila in ws.iter_rows(min_row=hdr + 1, values_only=True):
        if not any(fila):
            continue
        d = {c: fila[idx[c]] for c in idx}
        if d.get("codigo_indicador"):
            out.append(d)
    return out


def main():
    if not os.path.exists(BASE):
        sys.exit(f"No encuentro la base: {BASE}")
    wb = openpyxl.load_workbook(BASE, data_only=True)

    campos = ["periodo", "pais", "codigo_indicador", "segmento", "valor", "moneda"]
    bd = hoja_dict(wb["BD_Indicadores"], campos)

    # Las capas de soporte pueden estar en la base (antes de separar) o en el libro
    # de trazabilidad (después). Se buscan en ambos para que el control no se apague solo.
    wt = openpyxl.load_workbook(TRAZA, data_only=True) if os.path.exists(TRAZA) else None

    def capa(nombre):
        if nombre in wb.sheetnames:
            return hoja_dict(wb[nombre], campos), "base"
        if wt is not None and nombre in wt.sheetnames:
            return hoja_dict(wt[nombre], campos), "trazabilidad"
        return [], "NO ENCONTRADA"

    sub, donde_sub = capa("Subsidiarias_PyG")
    ofi, donde_ofi = capa("BD_PYG_OFICIAL")
    for nom, filas, donde in (("Subsidiarias_PyG", sub, donde_sub), ("BD_PYG_OFICIAL", ofi, donde_ofi)):
        if donde == "NO ENCONTRADA":
            print(f"  AVISO: no encuentro {nom} en ninguno de los dos libros — no se puede conciliar.")

    def val(filas, per, pais, cod, seg=None):
        for d in filas:
            if (str(d.get("periodo")) == per and d.get("pais") == pais
                    and d.get("codigo_indicador") == cod
                    and (seg is None or (d.get("segmento") or "") == seg)):
                return d.get("valor")
        return None

    meses = sorted({str(d["periodo"]) for d in bd if (d.get("segmento") or "") == SEG})
    problemas = []
    print("=" * 78)
    print("  CONCILIACIÓN DE CAPAS — BD_MAESTRA_COCO")
    print("=" * 78)
    print(f"  Meses del bloque consolidado: {', '.join(meses)}\n")

    # ---------- A. países vs consolidado ----------
    print("  A. PAÍSES vs CONSOLIDADO (bloque '%s')" % SEG)
    for per in meses:
        for cod in RUBROS:
            suma, faltan = 0.0, []
            for p in PAISES:
                v = val(bd, per, p, cod, SEG)
                if v is None:
                    faltan.append(p)
                elif isinstance(v, (int, float)):
                    suma += v
            cons = val(bd, per, "Consolidado", cod, SEG)
            if cons is None:
                continue
            dif = suma - cons
            if abs(dif) > TOL:
                problemas.append(("A", per, cod, f"suma países {suma:,.2f} vs consolidado {cons:,.2f} (dif {dif:,.2f})"))
                print(f"     {per}  {cod:30s} DIF {dif:>12,.2f}"
                      + (f"   [sin dato: {', '.join(faltan)}]" if faltan else ""))
    if not [p for p in problemas if p[0] == "A"]:
        print("     OK — todos los meses y rubros cuadran\n")
    else:
        print()

    # ---------- B y C. BD_Indicadores vs capas de soporte ----------
    # Solo son comparables cifras del MISMO segmento y la MISMA moneda. Sin ese filtro
    # se comparan peras con manzanas (Colombia en COP contra su equivalente en USD, o el
    # mes contra el acumulado del semestre) y el control se llena de falsos positivos.
    def busca(filas, per, pais, cod, seg, mon):
        for d in filas:
            if (str(d.get("periodo")) == per and d.get("pais") == pais
                    and d.get("codigo_indicador") == cod
                    and (d.get("segmento") or "") == seg
                    and (d.get("moneda") or "") == mon):
                return d.get("valor")
        return None

    for etiqueta, filas, letra, seg_fijo, donde in (("Subsidiarias_PyG", sub, "B", SEG, donde_sub),
                                                    ("BD_PYG_OFICIAL", ofi, "C", None, donde_ofi)):
        print(f"  {letra}. BD_Indicadores vs {etiqueta}   [en: {donde}]")
        if not filas:
            print("     (hoja ausente o sin datos legibles)\n")
            continue
        comparados = saltados = 0
        for d in filas:
            per, pais, cod = str(d.get("periodo")), d.get("pais"), d.get("codigo_indicador")
            v2, mon = d.get("valor"), (d.get("moneda") or "")
            if not isinstance(v2, (int, float)):
                continue
            # Subsidiarias_PyG no trae segmento: equivale al bloque mensual consolidado
            seg = seg_fijo if seg_fijo is not None else (d.get("segmento") or "")
            v1 = busca(bd, per, pais, cod, seg, mon)
            if v1 is None or not isinstance(v1, (int, float)):
                saltados += 1
                continue
            comparados += 1
            if abs(v1 - v2) > TOL:
                problemas.append((letra, per, f"{pais}/{cod}", f"BD {v1:,.2f} vs {etiqueta} {v2:,.2f}"))
                print(f"     {per}  {str(pais):12s} {cod:30s} BD {v1:>14,.2f} vs {v2:>14,.2f}  [{seg or 'mensual'}·{mon}]")
        hallados = len([p for p in problemas if p[0] == letra])
        print(f"     {comparados} cifras comparadas · {saltados} sin par en BD_Indicadores · "
              + ("OK — sin diferencias" if not hallados else f"{hallados} diferencia(s)") + "\n")

    # ---------- D. ceros sospechosos ----------
    print("  D. CEROS SOSPECHOSOS (posible dato no cargado, no operación real en cero)")
    sospechosos = 0
    for per in meses:
        for p in PAISES:
            vals = [val(bd, per, p, c, SEG) for c in RUBROS]
            presentes = [v for v in vals if isinstance(v, (int, float))]
            if presentes and all(v == 0 for v in presentes):
                sospechosos += 1
                problemas.append(("D", per, p, "todos los rubros en cero"))
                print(f"     {per}  {p}: todos los rubros en cero — confirmar si es real")
    if not sospechosos:
        print("     OK — ningún país-mes con todo en cero\n")
    else:
        print()

    # ---------- E. informe del Modelo ----------
    # La pestaña Burn & Runway no sale de la base: sale de burn_runway.json, que se
    # genera del Modelo. Es una copia derivada, y como toda copia derivada se queda
    # vieja sin avisar. Aquí se compara la huella del modelo con la que quedó sellada.
    print("  E. INFORME DEL MODELO (pestaña Burn & Runway)")
    modelo = os.path.join(AQUI, "..", "modelo", "Modelo_COCO.xlsx")
    informe = os.path.join(AQUI, "..", "burn_runway.json")
    if not os.path.exists(informe):
        print("     No existe burn_runway.json — la pestaña explica cómo generarlo.\n")
    elif not os.path.exists(modelo):
        print("     No encuentro modelo/Modelo_COCO.xlsx — no se puede verificar.\n")
    else:
        with open(informe, encoding="utf-8") as f:
            info = json.load(f)
        sello = info.get("fuente")
        if not sello:
            print("     El informe se generó antes de que existiera el sello de origen.")
            print("     Corre modelo/Actualizar_Burn_Runway.bat para poder verificarlo.\n")
            problemas.append(("E", info.get("generado", "?"), "burn_runway.json", "sin sello de origen"))
        else:
            h = hashlib.md5()
            with open(modelo, "rb") as f:
                for trozo in iter(lambda: f.read(1 << 20), b""):
                    h.update(trozo)
            actual = h.hexdigest()
            if actual == sello.get("md5"):
                print(f"     OK — el informe corresponde al modelo actual "
                      f"(corte {info.get('corte')}, generado {info.get('generado')})\n")
            else:
                problemas.append(("E", info.get("corte", "?"), "burn_runway.json",
                                  "el modelo cambió desde que se generó el informe"))
                print("     DESACTUALIZADO: el modelo cambió desde que se generó el informe.")
                print(f"       informe generado : {info.get('generado')}  (corte {info.get('corte')})")
                print(f"       modelo de entonces: {sello.get('bytes'):,} b   md5 {sello.get('md5')[:12]}")
                print(f"       modelo de ahora   : {os.path.getsize(modelo):,} b   md5 {actual[:12]}")
                print("       -> corre modelo/Actualizar_Burn_Runway.bat\n")

    # Aviso aparte: puede haber una versión más nueva del modelo sin copiar todavía.
    carpeta = os.path.join(AQUI, "..", "..", "Modelos Financieros")
    if os.path.isdir(carpeta) and os.path.exists(modelo):
        ref = os.path.getmtime(modelo)
        nuevos = [n for n in os.listdir(carpeta)
                  if n.lower().endswith(".xlsx") and "presupuesto completo" in n.lower()
                  and os.path.getmtime(os.path.join(carpeta, n)) > ref]
        if nuevos:
            print("     AVISO: hay modelo(s) más recientes que la copia del tablero:")
            for n in sorted(nuevos)[-3:]:
                print(f"       {n}")
            print("     Confirma cuál es el vigente antes de reemplazar la copia.\n")

    # ---------- resumen ----------
    print("=" * 78)
    if problemas:
        print(f"  RESULTADO: {len(problemas)} punto(s) a revisar.")
        print("  Corrige en la capa que corresponda y vuelve a ejecutar antes de publicar.")
    else:
        print("  RESULTADO: todas las capas están alineadas. Base lista para publicar.")
    print("=" * 78)
    return 1 if problemas else 0


if __name__ == "__main__":
    sys.exit(main())
