# -*- coding: utf-8 -*-
"""
Agrega el P&G Consolidado OFICIAL — Acumulado a Junio 2026 (H1, USD) a tbl_BD_Indicadores.

Fuente: estado financiero consolidado entregado por Finanzas (1 ene - 30 jun 2026,
sin balance general). Se guarda como escenario='Real', segmento='Acumulado H1 2026'
(igual convención que 'Original enero'/'Revisado mayo': un vintage/corte explícito
que NO se suma por defecto junto con los hechos mensuales — se consulta aparte).

Verificado antes de escribir: subsidiarias suman exacto al consolidado (excepto
redondeos de 1 USD), y Resultado de la operación / Resultado antes de impuestos
cuadran con ingresos-costo-gastos, igual que la fórmula que ya usa el dashboard
para el EBITDA consolidado mensual.
"""
import openpyxl

RUTA = r"C:\Users\andre\OneDrive - CFOcus\01_CFOCUS_CLIENTES\COCO\1.0 Coco Digital\COCO_Dashboard_Cloud\migracion\BD_COCO_2026.xlsx"
HOJA = "BD_Indicadores"
TABLA = "tbl_BD_Indicadores"
PERIODO = "2026-06"
SEGMENTO = "Acumulado H1 2026"

PAISES = {
    "EE.UU.":     "Coco LLC",
    "Perú":       "Coco Perú",
    "Costa Rica": "Coco Costa Rica",
    "Colombia":   "Coco Colombia",
    "Consolidado":"Coco Consolidado",
}

# codigo_indicador -> {pais: valor}  (None = "-" en el estado fuente, no se carga)
DATOS = {
    "pyg_servicio_software":        {"EE.UU.":32997, "Perú":4679,  "Costa Rica":None,  "Colombia":1398926, "Consolidado":1436602},
    "pyg_devoluciones":             {"EE.UU.":None,  "Perú":None,  "Costa Rica":None,  "Colombia":-123034, "Consolidado":-123034},
    "pyg_ingresos_operacionales":   {"EE.UU.":32997, "Perú":4679,  "Costa Rica":None,  "Colombia":1275892, "Consolidado":1313567},
    "pyg_costo_ventas":             {"EE.UU.":None,  "Perú":None,  "Costa Rica":None,  "Colombia":605005,  "Consolidado":605005},
    "pyg_gasto_administracion":     {"EE.UU.":30736, "Perú":1600,  "Costa Rica":21792, "Colombia":196717,  "Consolidado":250845},
    "pyg_gasto_ventas":             {"EE.UU.":25216, "Perú":None,  "Costa Rica":None,  "Colombia":240203,  "Consolidado":265419},
    "pyg_gasto_proyectos":          {"EE.UU.":11500, "Perú":16797, "Costa Rica":None,  "Colombia":163589,  "Consolidado":191885},
    "pyg_ingresos_no_operacionales":{"EE.UU.":None,  "Perú":None,  "Costa Rica":None,  "Colombia":2876,    "Consolidado":2876},
    "pyg_gastos_financieros":       {"EE.UU.":15,    "Perú":265,   "Costa Rica":206,   "Colombia":9915,    "Consolidado":10401},
    "pyg_utilidad_neta":            {"EE.UU.":-34470,"Perú":-13983,"Costa Rica":-21998,"Colombia":63339,   "Consolidado":-7112},
}

COMENTARIO = ("P&G Consolidado Oficial · Acumulado 1-ene al 30-jun-2026 (H1) · USD · sin balance general. "
              "Tasas promedio jun-2026: Perú 3.4050 PEN/USD · Colombia 3,505.2683 COP/USD.")


def fila_base(pais):
    return {"pais":pais,"compania":PAISES[pais],"moneda":"USD","escenario":"Real",
            "segmento":SEGMENTO,"detalle":"","unidad":"USD","comentario":COMENTARIO,
            "area":"Internacional","area_responsable":"Finanzas","periodicidad":"Semestral",
            "periodo":PERIODO}


def construir_filas():
    filas = []
    for codigo, porpais in DATOS.items():
        for pais, valor in porpais.items():
            if valor is None:
                continue  # "-" en el estado fuente: no se carga (evita inventar ceros)
            f = fila_base(pais)
            f["codigo_indicador"] = codigo
            f["valor"] = valor
            filas.append(f)
    return filas


def main():
    wb = openpyxl.load_workbook(RUTA)
    ws = wb[HOJA]
    tabla = ws.tables[TABLA]
    encabezados = [c.value for c in ws[3]]

    fila_inicio = ws.max_row + 1
    filas = construir_filas()
    print(f"Filas a insertar: {len(filas)} (desde la fila {fila_inicio})")

    for i, f in enumerate(filas):
        r = fila_inicio + i
        for col_idx, nombre_col in enumerate(encabezados, start=1):
            ws.cell(row=r, column=col_idx, value=f.get(nombre_col, ""))

    fila_fin = fila_inicio + len(filas) - 1
    col_ini, col_fin_actual = tabla.ref.split(":")
    col_letra_ini = ''.join(ch for ch in col_ini if ch.isalpha())
    fila_ini_tabla = ''.join(ch for ch in col_ini if ch.isdigit())
    col_letra_fin = ''.join(ch for ch in col_fin_actual if ch.isalpha())
    nuevo_ref = f"{col_letra_ini}{fila_ini_tabla}:{col_letra_fin}{fila_fin}"
    print("Ref de tabla:", tabla.ref, "->", nuevo_ref)
    tabla.ref = nuevo_ref

    wb.save(RUTA)
    print(f"\nListo. {len(filas)} filas del H1 2026 oficial agregadas a {TABLA}.")


if __name__ == "__main__":
    main()
