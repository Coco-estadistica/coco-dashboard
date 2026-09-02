# -*- coding: utf-8 -*-
"""
Carga el saldo de caja de julio 2026 para Liquidez & Cartera (codigo liquidez_total),
que nunca se habia cargado para ningun pais ni periodo -- la tarjeta del tablero decia
"Por medir" desde siempre.

Fuente: reportado directamente por el usuario en la sesion del 31-ago-2026, confirmado
que corresponde al cierre de julio 2026 (no es un corte del dia). No es un PDF ni un
EEFF oficial -- es una cifra de tesoreria, consistente con lo que pide el Diccionario
para este codigo (Responsable carga: Tesoreria; Fuente sugerida: Bancos).

  Estados Unidos:  USD 276.523      (moneda nativa: USD)
  Colombia:        COP 251.790.666  (moneda nativa: COP)

Por que dos filas y no una consolidada
---------------------------------------
Cada pais entra en su moneda de origen, igual que el resto de la base (regla del
proyecto: no convertir a mano antes de cargar). El tablero las convierte y las suma
el solo -- pero para que sepa sumar USD y COP en el mismo total, rawStored() en
Dashboard_COCO.html tenia un vacio: nunca miraba el campo "moneda" de la fila, asumia
que todo ya estaba en COP. Se corrigio esa funcion (aCopSiHaceFalta) para que convierta
a COP antes de sumar cuando la fila dice USD o PEN. Sin ese cambio, la fila de Estados
Unidos se hubiera sumado como si 276.523 fueran pesos.

Con las dos filas en su moneda nativa y ese arreglo, la tarjeta de Liquidez Total:
  - bajo pais=Colombia solo, muestra la fila de Colombia (conversion normal via TRM).
  - bajo pais=Consolidado, muestra la SUMA de ambas, ya convertida, y responde al
    selector de moneda (USD/COP) igual que cualquier otro indicador.
  - deja de mostrar el chip "solo Colombia": ya hay mas de un pais reportando.

Uso:  python cargar_liquidez_julio_2026.py             -> vista previa
      python cargar_liquidez_julio_2026.py --escribir  -> aplica (con respaldo)
"""
import os
import shutil
import sys
from datetime import datetime

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.join(AQUI, "BD_MAESTRA_COCO.xlsx")
HOJA = "BD_Indicadores"
PERIODO = "2026-07"
CODIGO = "liquidez_total"
COMENTARIO = "Saldo bancario reportado por el usuario (sesion 31-ago-2026); corresponde al cierre de julio 2026."

FILAS = [
    # periodo, pais, compania, moneda, escenario, codigo, segmento, detalle, valor, unidad, comentario, area, area_responsable, periodicidad
    (PERIODO, "EE.UU.", "Coco LLC", "USD", "Real", CODIGO, None, "Total bancos", 276523, "USD", COMENTARIO, "Bancos", "Tesorería", "Mensual"),
    (PERIODO, "Colombia", "Coco Colombia", "COP", "Real", CODIGO, None, "Total bancos", 251790666, "COP", COMENTARIO, "Bancos", "Tesorería", "Mensual"),
]

COLUMNAS = ["periodo", "pais", "compania", "moneda", "escenario", "codigo_indicador",
            "segmento", "detalle", "valor", "unidad", "comentario", "area", "area_responsable", "periodicidad"]


def main():
    escribir = "--escribir" in sys.argv
    wb = openpyxl.load_workbook(BASE)
    ws = wb[HOJA]
    H = [str(c.value).strip() if c.value else "" for c in ws[3]]
    ix = {h: i + 1 for i, h in enumerate(H)}  # 1-based para openpyxl

    # ya existe alguna fila de liquidez_total para julio? (evita duplicar si se corre dos veces)
    cPer, cCod, cPais = ix["periodo"], ix["codigo_indicador"], ix["pais"]
    existentes = []
    for r in range(4, ws.max_row + 1):
        if ws.cell(row=r, column=cCod).value == CODIGO and str(ws.cell(row=r, column=cPer).value) == PERIODO:
            existentes.append((r, ws.cell(row=r, column=cPais).value))

    print("Filas a insertar:")
    for fila in FILAS:
        d = dict(zip(COLUMNAS, fila))
        print("  %-9s %-9s %-9s %14s %s" % (d["periodo"], d["pais"], d["moneda"],
              "{:,.2f}".format(d["valor"]), d["comentario"][:50]))

    if existentes:
        print("\n!! ADVERTENCIA: ya hay %d fila(s) de %s en %s:" % (len(existentes), CODIGO, PERIODO))
        for r, p in existentes:
            print("   fila %d, pais=%s" % (r, p))
        print("   No se insertara de nuevo. Revisa a mano si hace falta corregir en vez de agregar.")
        return

    trm_julio = 3268.95
    usd_col = FILAS[1][8] / trm_julio
    total_usd = FILAS[0][8] + usd_col
    total_cop = FILAS[0][8] * trm_julio + FILAS[1][8]
    print("\nVerificacion del consolidado (a TRM julio %.2f):" % trm_julio)
    print("  Colombia en USD: %s" % "{:,.2f}".format(usd_col))
    print("  TOTAL consolidado: USD %s  |  COP %s" % ("{:,.2f}".format(total_usd), "{:,.2f}".format(total_cop)))

    if not escribir:
        print("\nVISTA PREVIA. Nada se escribio. Corre con --escribir para aplicar.")
        return

    nueva_fila = ws.max_row + 1
    for fila in FILAS:
        for col_name, valor in zip(COLUMNAS, fila):
            ws.cell(row=nueva_fila, column=ix[col_name]).value = valor
        nueva_fila += 1

    os.makedirs(os.path.join(AQUI, "respaldos"), exist_ok=True)
    sello = datetime.now().strftime("%Y%m%d_%H%M%S")
    copia = os.path.join(AQUI, "respaldos", "BD_MAESTRA_COCO_antes_liquidez_julio_" + sello + ".xlsx")
    shutil.copy2(BASE, copia)
    print("\nrespaldo -> %s" % os.path.basename(copia))
    wb.save(BASE)
    print("ESCRITO  -> %s (2 filas nuevas)" % os.path.basename(BASE))


if __name__ == "__main__":
    main()
