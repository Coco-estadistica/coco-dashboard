# -*- coding: utf-8 -*-
"""
Recarga el P&G mensual (segmento "Consolidacion USD") de EE.UU./Peru/Costa Rica
para enero-julio 2026, con la contabilidad como unico criterio (auxiliares
homologados y validados: cada mes real, sin duplicados, encadenando saldo a
saldo). Reemplaza lo que hoy sale de una mezcla de modelo financiero (abr-jul)
y nada (ene-mar).

Que SI hace:
  - Escribe pyg_* de EE.UU./Peru/Costa Rica, ene-jul 2026.
  - Recalcula Consolidado (segmento Consolidacion USD) para abr-jul, sumando
    los 4 paises (Colombia + las 3 filiales).

Que NO hace (a proposito):
  - No toca Colombia. Su P&G propio (ene-mar en COP, abr-jul en
    Consolidacion USD) no sale de este proceso.
  - No crea Consolidado para ene-mar: el ingreso de Colombia ene-mar guardado
    en la base es de gestion (no concilia con el PDF oficial: la suma da
    USD 514.030 contra los USD 630.459 esperados por la reconciliacion H1).
    Convertirlo a USD y sumarlo repetiria el mismo problema en otro lugar.
    Ese hueco queda abierto, documentado, no relleno con datos de otra calidad.
  - No toca pyg_reclasif_honorarios_peru (vive solo en el segmento
    "Acumulado H1 2026", fuera de este proceso).

Uso:  python cargar_pyg_filiales_ene_jul.py <ruta_json>            -> vista previa
      python cargar_pyg_filiales_ene_jul.py <ruta_json> --escribir -> aplica (con respaldo)
"""
import json
import os
import shutil
import sys
from datetime import datetime

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.join(AQUI, "BD_MAESTRA_COCO.xlsx")
CODES = ["pyg_ingresos_operacionales", "pyg_ingresos_no_operacionales", "pyg_costo_ventas",
          "pyg_gasto_administracion", "pyg_gasto_ventas", "pyg_gasto_proyectos",
          "pyg_gastos_financieros", "pyg_utilidad_neta"]
COMPANIA = {"EE.UU.": "Coco LLC", "Perú": "Coco Peru", "Costa Rica": "Coco Costa Rica"}
# OJO: el JSON generado por extraer_pyg_filiales.py usa "Perú" CON tilde como clave
# (asi quedo definido su propio diccionario MEJOR). Si aqui se itera con "Peru" sin
# tilde, nuevo.get(...) nunca encuentra la clave y el pais se salta ENTERO sin avisar
# -- exactamente lo que paso la primera vez que corrio este script.
PAIS_JSON_A_BASE = {"EE.UU.": "EE.UU.", "Perú": "Perú", "Costa Rica": "Costa Rica"}
PERIODOS = ["2026-01", "2026-02", "2026-03", "2026-04", "2026-05", "2026-06", "2026-07"]


def main(ruta_json, escribir):
    nuevo = json.load(open(ruta_json, encoding="utf-8"))

    wb = openpyxl.load_workbook(BASE)
    ws = wb["BD_Indicadores"]
    H = [str(c.value).strip() if c.value else "" for c in ws[3]]
    ix = {c: H.index(c) + 1 for c in H if c}

    existentes = {}
    for f in range(4, ws.max_row + 1):
        pais = ws.cell(f, ix["pais"]).value
        per = ws.cell(f, ix["periodo"]).value
        cod = ws.cell(f, ix["codigo_indicador"]).value
        seg = ws.cell(f, ix["segmento"]).value or ""
        esc = ws.cell(f, ix["escenario"]).value
        if seg == "Consolidación USD" and esc == "Real" and cod:
            existentes[(pais, str(per), cod)] = f

    cambios, nuevas = 0, 0
    reporte = []

    def escribe_fila(fila, pais, per, cod, valor, comentario, compania):
        ws.cell(fila, ix["periodo"]).value = per
        ws.cell(fila, ix["pais"]).value = pais
        ws.cell(fila, ix["compania"]).value = compania
        ws.cell(fila, ix["moneda"]).value = "USD"
        ws.cell(fila, ix["escenario"]).value = "Real"
        ws.cell(fila, ix["codigo_indicador"]).value = cod
        ws.cell(fila, ix["segmento"]).value = "Consolidación USD"
        ws.cell(fila, ix["valor"]).value = valor
        ws.cell(fila, ix["unidad"]).value = "USD"
        ws.cell(fila, ix["comentario"]).value = comentario
        ws.cell(fila, ix["area"]).value = "Internacional"
        ws.cell(fila, ix["area_responsable"]).value = "Finanzas"
        ws.cell(fila, ix["periodicidad"]).value = "Mensual"

    siguiente_fila = ws.max_row + 1

    for pais_json in ["EE.UU.", "Perú", "Costa Rica"]:
        pais = PAIS_JSON_A_BASE[pais_json]
        for per in PERIODOS:
            datos = nuevo.get(pais_json + "|" + per)
            if not datos:
                continue
            fuente = datos.get("_fuente") or "?"
            comentario = "Homologado desde auxiliar contable (" + fuente + "). Reemplaza fuente previa (modelo/sin dato)."
            for cod in CODES:
                valor = datos[cod]
                clave = (pais, per, cod)
                if clave in existentes:
                    f = existentes[clave]
                    antes = ws.cell(f, ix["valor"]).value
                    if antes is None or abs((antes or 0) - valor) > 0.01:
                        cambios += 1
                        reporte.append("  UPDATE %-12s %s %-30s %s -> %s" % (pais, per, cod, antes, valor))
                    escribe_fila(f, pais, per, cod, valor, comentario, COMPANIA[pais_json])
                else:
                    escribe_fila(siguiente_fila, pais, per, cod, valor, comentario, COMPANIA[pais_json])
                    existentes[clave] = siguiente_fila
                    nuevas += 1
                    reporte.append("  INSERT %-12s %s %-30s -> %s" % (pais, per, cod, valor))
                    siguiente_fila += 1

    consol_cambios = []
    for per in ["2026-04", "2026-05", "2026-06", "2026-07"]:
        for cod in CODES:
            suma = 0.0
            falta = []
            for pais in ["Colombia", "EE.UU.", "Perú", "Costa Rica"]:
                clave = (pais, per, cod)
                if clave in existentes:
                    v = ws.cell(existentes[clave], ix["valor"]).value
                    suma += v or 0
                else:
                    falta.append(pais)
            if falta:
                consol_cambios.append("  AVISO Consolidado %s %s: falta %s, no se recalcula" % (per, cod, falta))
                continue
            clave = ("Consolidado", per, cod)
            comentario = "Recalculado: suma Colombia + EE.UU. + Peru + Costa Rica (Consolidacion USD)."
            if clave in existentes:
                f = existentes[clave]
                antes = ws.cell(f, ix["valor"]).value
                if antes is None or abs((antes or 0) - round(suma, 2)) > 0.01:
                    consol_cambios.append("  UPDATE Consolidado %s %-30s %s -> %s" % (per, cod, antes, round(suma, 2)))
                escribe_fila(f, "Consolidado", per, cod, round(suma, 2), comentario, "COCO Tecnologias")
            else:
                escribe_fila(siguiente_fila, "Consolidado", per, cod, round(suma, 2), comentario, "COCO Tecnologias")
                existentes[clave] = siguiente_fila
                consol_cambios.append("  INSERT Consolidado %s %-30s -> %s" % (per, cod, round(suma, 2)))
                siguiente_fila += 1

    print("=" * 90)
    print("  RECARGA P&G FILIALES (contabilidad) -- ene-jul 2026")
    print("=" * 90)
    print("\n  Filiales: %d valores cambian, %d filas nuevas (de %d posibles)\n" % (cambios, nuevas, len(CODES) * 3 * 7))
    for l in reporte[:60]:
        print(l)
    if len(reporte) > 60:
        print("  ... y %d mas" % (len(reporte) - 60))
    print("\n  Consolidado (abr-jul, recalculado por suma):")
    for l in consol_cambios:
        print(l)
    print("\n  Consolidado ene-mar: NO se toca (Colombia no tiene bloque oficial en USD para esos meses).")

    if not escribir:
        print("\n  Vista previa. Para aplicar agrega --escribir\n")
        return

    os.makedirs(os.path.join(AQUI, "respaldos"), exist_ok=True)
    sello = datetime.now().strftime("%Y%m%d_%H%M%S")
    copia = os.path.join(AQUI, "respaldos", "BD_MAESTRA_COCO_antes_pyg_filiales_" + sello + ".xlsx")
    shutil.copy2(BASE, copia)
    wb.save(BASE)
    print("\n  Respaldo: " + os.path.basename(copia))
    print("  Base actualizada: " + BASE)
    print("  Siguiente: corre Revisar_Base.bat y abre el tablero para confirmar.\n")


if __name__ == "__main__":
    ruta = sys.argv[1]
    escribir = "--escribir" in sys.argv
    main(ruta, escribir)
