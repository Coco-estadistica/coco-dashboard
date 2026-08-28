# -*- coding: utf-8 -*-
"""
Borra el indicador "ebitda" viejo (Colombia, COP, sin comentario/fuente) que
quedaba como respaldo para los meses sin P&G oficial. Se decidio borrarlo
en vez de dejarlo como respaldo silencioso: es una cifra sin trazabilidad
(sin comentario, sin desglose, no verificada contra el ERP) y da una serie
sospechosamente pareja (ago25-mar26: 46-54K USD cada mes) frente a lo volatil
que es el P&G real ya reconstruido (abr-jul 2026).

Cubre 13 filas, jun-2025 a jun-2026 (abr-jun-2026 ya estaban de cualquier
forma superadas por el calculo oficial via pyg_*; jun25-mar26 quedan sin
dato en vez de una estimacion sin trazabilidad -- mas honesto que un
numero que parece preciso pero no se puede verificar).

Uso:  python borrar_ebitda_legacy.py             -> vista previa
      python borrar_ebitda_legacy.py --escribir  -> aplica (con respaldo)
"""
import os
import shutil
import sys
from datetime import datetime

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.join(AQUI, "BD_MAESTRA_COCO.xlsx")


def main():
    escribir = "--escribir" in sys.argv
    wb = openpyxl.load_workbook(BASE)
    ws = wb["BD_Indicadores"]
    H = [str(c.value).strip() if c.value else "" for c in ws[3]]
    ix = {c: H.index(c) + 1 for c in H if c}

    filas_borrar = []
    for f in range(4, ws.max_row + 1):
        if ws.cell(f, ix["codigo_indicador"]).value == "ebitda":
            filas_borrar.append((f, ws.cell(f, ix["periodo"]).value, ws.cell(f, ix["valor"]).value))

    print("=" * 90)
    print("  BORRAR ebitda (legacy, Colombia COP, sin fuente) -- %d filas" % len(filas_borrar))
    print("=" * 90)
    print()
    for f, per, val in filas_borrar:
        print("  DELETE  fila %-6d %-10s -> %s" % (f, per, val))

    if not escribir:
        print("\n  Vista previa. Para aplicar agrega --escribir\n")
        return

    os.makedirs(os.path.join(AQUI, "respaldos"), exist_ok=True)
    sello = datetime.now().strftime("%Y%m%d_%H%M%S")
    copia = os.path.join(AQUI, "respaldos", "BD_MAESTRA_COCO_antes_borrar_ebitda_legacy_" + sello + ".xlsx")
    shutil.copy2(BASE, copia)

    # borra de abajo hacia arriba para no correr los indices de fila
    for f, per, val in sorted(filas_borrar, key=lambda x: -x[0]):
        ws.delete_rows(f, 1)

    wb.save(BASE)
    print("\n  Respaldo:", os.path.basename(copia))
    print("  Base actualizada. %d filas borradas.\n" % len(filas_borrar))


if __name__ == "__main__":
    main()
