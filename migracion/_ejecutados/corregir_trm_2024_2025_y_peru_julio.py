# -*- coding: utf-8 -*-
"""
Corrige tres tasas de cambio que quedaron con valores de relleno o parciales.
Todas comparten la misma causa raiz: se capturaron a mitad de periodo (o se
pusieron fijas "por ahora") y nunca se sustituyeron por el dato definitivo.

--------------------------------------------------------------------------
1. TRM Colombia ene-24 a dic-25: 4.420 FIJO los 24 meses  ->  TRM de mercado
--------------------------------------------------------------------------
Es exactamente el mismo defecto que corregir_trm_ene_jun.py arreglo para 2026,
pero dos anios completos. Veinticuatro meses a la misma tasa al peso no es una
tasa implicita: es relleno.

La distorsion NO es pareja, y ese es el problema. Va de -0,3% (nov-24, cuando
4.420 casi coincide con el mercado) a -14,2% (dic-25, cuando el mercado estaba
en 3.791). O sea que cualquier serie en COP mostrada en USD durante 2024-2025
trae crecimiento inventado por efecto cambiario, igual que el MRR de julio que
aparecia creciendo 52,3% cuando en pesos crecia 19,4%.

De donde sale la tasa nueva: de la columna trm_promedio_mercado de la MISMA
hoja, que no se toca. Se verifico contra la TRM oficial diaria publicada en
datos.gov.co (Superfinanciera), promediada por DIAS CALENDARIO -- cada registro
de TRM rige de vigenciadesde a vigenciahasta, y los viernes cubren sabado y
domingo. Con esa convencion la columna reproduce la fuente oficial al centavo
en los 24 meses. No hay que reemplazarla; hay que usarla.

Por que mercado y no derivada: 2024-2025 no tiene bloque de consolidacion en
USD de Colombia, asi que no hay nada de donde derivar. Es el mismo criterio ya
aplicado a ene-mar 2026.

Seguridad del cambio: BD_Indicadores no tiene NI UN registro en USD en 2024 ni
en 2025 (417 y 499 registros, todos COP; los 68 sin moneda son %, conteos,
veces y meses). El tablero convierte al vuelo con trm[periodo], no hay valores
USD almacenados que queden descuadrados. Por eso este script NO reconvierte
nada, a diferencia de corregir_trm_julio.py, que si tuvo que reescribir los
pyg_* de julio porque 2026 si tiene USD guardado.

--------------------------------------------------------------------------
2. trm_promedio_mercado de jun-26: 3.571,44  ->  3.500,98
--------------------------------------------------------------------------
3.571,44 es el promedio de junio 1 a 13 UNICAMENTE (verificado: esos 13 dias
promedian 3.571,4423). El mes completo, ponderado por dias calendario, da
3.500,98. Es el unico mes de toda la serie que no calza con la fuente oficial.

No se toca trm_cop_usd de jun-26 (3.505,27): esa es la implicita derivada del
bloque de consolidacion (751.453.314 / 214.378,26) y es la que hace cuadrar el
P&G. Al corregir el mercado, la brecha implicita/mercado de junio pasa de
-1,85% a +0,12%, que es lo razonable.

--------------------------------------------------------------------------
3. trm_promedio_mercado de jul-26: vacio  ->  3.268,95
--------------------------------------------------------------------------
Faltaba. Con la celda vacia, al elegir «Base TRM: Mercado» el tablero se queda
sin tasa para julio y muestra "Sin las dos tasas para este periodo". El
promedio calendario de julio da 3.268,95, identico a la tasa oficial que ya
esta en trm_cop_usd.

--------------------------------------------------------------------------
4. TRM_Peru jul-26: 3,405 provisional  ->  3,400 definitivo
--------------------------------------------------------------------------
La hoja traia "2026-07*" con nota "promedio hasta 9-jul (provisional)", y el
valor 3,405 era simplemente el de junio repetido.

Dato definitivo del BCRP, serie PN01246PM "Tipo de cambio - promedio del
periodo (S/ por US$)": julio 2026 = 3,400. Contrastado ademas contra las series
diarias del sistema bancario SBS de julio: compra promedia 3,395474 y venta
3,404526, cuyo punto medio da exactamente 3,400000. Enero a junio de la serie
mensual calzan al milesimo con lo que ya estaba en la hoja, lo que confirma que
es la misma serie que se venia usando.

factor_usd_pen = 1/3,400 = 0,294118 (6 decimales, igual que el resto).

Uso:  python corregir_trm_2024_2025_y_peru_julio.py             -> vista previa
      python corregir_trm_2024_2025_y_peru_julio.py --escribir  -> aplica (con respaldo)
"""
import os
import shutil
import sys
from datetime import datetime

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.join(AQUI, "BD_MAESTRA_COCO.xlsx")

TRM_RELLENO = 4420.0
FUENTE_MERCADO = "TRM promedio de mercado (no derivable: sin bloque USD de Colombia)"

# Promedios calendario oficiales (datos.gov.co / Superfinanciera), verificados.
MERCADO_FIX = {"2026-06": 3500.98, "2026-07": 3268.95}

PERU_PERIODO_VIEJO = "2026-07*"
PERU_PERIODO_NUEVO = "2026-07"
PERU_TC = 3.400
PERU_FACTOR = round(1 / 3.400, 6)
PERU_FUENTE = "BCRP promedio mensual"
PERU_NOTA = ("Fuente: BCRP (promedio del periodo). Serie PN01246PM; julio 2026 "
             "definitivo (3,400 = punto medio compra/venta SBS del mes completo).")


def encabezado(ws, maxfilas=12):
    """Fila cuyo primer campo es exactamente 'periodo'."""
    for r in range(1, min(ws.max_row, maxfilas) + 1):
        v = ws.cell(row=r, column=1).value
        if v and str(v).strip().lower() == "periodo":
            return r
    raise SystemExit("No encontre la fila de encabezado en la hoja %s" % ws.title)


def cols(ws, fila):
    m = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=fila, column=c).value
        if v:
            m[str(v).strip().lower()] = c
    return m


def fmt(v):
    if v is None:
        return "(vacio)"
    if isinstance(v, float) and v < 10:
        return "%.6f" % v
    if isinstance(v, (int, float)):
        return "{:,.2f}".format(v)
    return str(v)[:26]


def main():
    escribir = "--escribir" in sys.argv
    wb = openpyxl.load_workbook(BASE)
    cambios = []

    # ---------------- Hoja TRM ----------------
    ws = wb["TRM"]
    h = encabezado(ws)
    cm = cols(ws, h)
    cPer, cCop, cFte, cMer = cm["periodo"], cm["trm_cop_usd"], cm["fuente"], cm["trm_promedio_mercado"]

    for r in range(h + 1, ws.max_row + 1):
        per = ws.cell(row=r, column=cPer).value
        if not per:
            continue
        per = str(per).strip()
        cop = ws.cell(row=r, column=cCop).value
        mer = ws.cell(row=r, column=cMer).value

        # 1. relleno 4420 -> mercado
        if per[:4] in ("2024", "2025") and cop is not None and abs(float(cop) - TRM_RELLENO) < 1e-6:
            if mer is None:
                print("  !! %s no tiene TRM de mercado; se deja como esta" % per)
                continue
            nuevo = round(float(mer), 2)
            cambios.append(("TRM", r, per, "trm_cop_usd", cop, nuevo))
            if escribir:
                ws.cell(row=r, column=cCop).value = nuevo
                ws.cell(row=r, column=cFte).value = FUENTE_MERCADO

        # 2 y 3. mercado parcial o vacio
        if per in MERCADO_FIX:
            nuevo = MERCADO_FIX[per]
            actual = None if mer is None else round(float(mer), 2)
            if actual != nuevo:
                cambios.append(("TRM", r, per, "trm_promedio_mercado", actual, nuevo))
                if escribir:
                    ws.cell(row=r, column=cMer).value = nuevo

    # ---------------- Hoja TRM_Peru ----------------
    wp = wb["TRM_Peru"]
    hp = encabezado(wp)
    pm = cols(wp, hp)
    pPer, pTc, pFac, pFte = pm["periodo"], pm["tc_pen_usd"], pm["factor_usd_pen"], pm["fuente"]

    encontrado = False
    for r in range(hp + 1, wp.max_row + 1):
        per = wp.cell(row=r, column=pPer).value
        if not per or str(per).strip() != PERU_PERIODO_VIEJO:
            continue
        encontrado = True
        cambios.append(("TRM_Peru", r, PERU_PERIODO_NUEVO, "periodo",
                        PERU_PERIODO_VIEJO, PERU_PERIODO_NUEVO))
        cambios.append(("TRM_Peru", r, PERU_PERIODO_NUEVO, "tc_pen_usd",
                        wp.cell(row=r, column=pTc).value, PERU_TC))
        cambios.append(("TRM_Peru", r, PERU_PERIODO_NUEVO, "factor_usd_pen",
                        wp.cell(row=r, column=pFac).value, PERU_FACTOR))
        if escribir:
            wp.cell(row=r, column=pPer).value = PERU_PERIODO_NUEVO
            wp.cell(row=r, column=pTc).value = PERU_TC
            wp.cell(row=r, column=pFac).value = PERU_FACTOR
            wp.cell(row=r, column=pFte).value = PERU_FUENTE
    if not encontrado:
        print("  !! no encontre la fila '%s' en TRM_Peru (ya corregida?)" % PERU_PERIODO_VIEJO)

    # nota de encabezado que decia "provisional"
    nota = wp.cell(row=2, column=1).value
    if nota and "provisional" in str(nota).lower():
        cambios.append(("TRM_Peru", 2, "-", "nota encabezado", "...provisional...", "...definitivo..."))
        if escribir:
            wp.cell(row=2, column=1).value = PERU_NOTA

    # ---------------- salida ----------------
    print("")
    print("%-9s %-5s %-9s %-22s %14s    %14s" % ("HOJA", "FILA", "PERIODO", "CAMPO", "ANTES", "DESPUES"))
    print("-" * 94)
    for hoja, fila, per, campo, antes, desp in cambios:
        print("%-9s %-5d %-9s %-22s %14s -> %14s" % (hoja, fila, per, campo, fmt(antes), fmt(desp)))
    print("-" * 94)
    print("TOTAL: %d cambios" % len(cambios))

    if not escribir:
        print("")
        print("VISTA PREVIA. Nada se escribio. Corre con --escribir para aplicar.")
        return

    os.makedirs(os.path.join(AQUI, "respaldos"), exist_ok=True)
    sello = datetime.now().strftime("%Y%m%d_%H%M%S")
    copia = os.path.join(AQUI, "respaldos",
                         "BD_MAESTRA_COCO_antes_trm_2024_2025_peru_" + sello + ".xlsx")
    shutil.copy2(BASE, copia)
    print("")
    print("respaldo -> %s" % os.path.basename(copia))
    wb.save(BASE)
    print("ESCRITO  -> %s" % os.path.basename(BASE))


if __name__ == "__main__":
    main()
