# -*- coding: utf-8 -*-
"""
Carga junio 2026 de "costo_proveedor_ventas_pct" (Colombia). Formula validada
contra enero-mayo ya cargados (costo_proveedor total / ingresos operacionales
oficiales, ambos Colombia COP) -- coincide dentro de +/-1pp con lo ya
guardado en esos 5 meses:

  costo_proveedor junio (188,300,000) / ingresos_operacionales junio
  (751,453,314) = 25.06%

Julio NO se carga aqui: costo_proveedor de julio esta incompleto (solo
Infobip+B2Chat, 2 de 5 proveedores -- falta AWS/IVA/Otros), asi que el
porcentaje saldria artificialmente bajo. Se carga cuando se resuelva ese
hueco.

Uso:  python cargar_proveedor_ventas_pct_junio.py             -> vista previa
      python cargar_proveedor_ventas_pct_junio.py --escribir  -> aplica
"""
import os
import shutil
import sys
from datetime import datetime

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.join(AQUI, "BD_MAESTRA_COCO.xlsx")
PERIODO = "2026-06"
VALOR = 25.06
COMENTARIO = ("Junio 2026 = costo_proveedor total (188.3M) / ingresos_operacionales "
              "oficiales (751.45M), Colombia COP. Formula validada contra ene-may ya "
              "cargados. Julio pendiente: costo_proveedor aun incompleto (falta AWS/IVA/Otros).")


def main():
    escribir = "--escribir" in sys.argv
    wb = openpyxl.load_workbook(BASE)
    ws = wb["BD_Indicadores"]
    H = [str(c.value).strip() if c.value else "" for c in ws[3]]
    ix = {c: H.index(c) + 1 for c in H if c}

    f = None
    for r in range(4, ws.max_row + 1):
        if (ws.cell(r, ix["codigo_indicador"]).value == "costo_proveedor_ventas_pct"
                and ws.cell(r, ix["pais"]).value == "Colombia"
                and str(ws.cell(r, ix["periodo"]).value) == PERIODO):
            f = r
            break

    print("=" * 90)
    print("  CARGA costo_proveedor_ventas_pct - JUNIO 2026 (Colombia)")
    print("=" * 90)

    if f:
        antes = ws.cell(f, ix["valor"]).value
        print("\n  UPDATE  %s -> %s" % (antes, VALOR))
    else:
        f = ws.max_row + 1
        ws.cell(f, ix["periodo"]).value = PERIODO
        ws.cell(f, ix["pais"]).value = "Colombia"
        ws.cell(f, ix["compania"]).value = "Coco Colombia"
        ws.cell(f, ix["escenario"]).value = "Real"
        ws.cell(f, ix["periodicidad"]).value = "Mensual"
        print("\n  INSERT  -> %s" % VALOR)
    ws.cell(f, ix["codigo_indicador"]).value = "costo_proveedor_ventas_pct"
    ws.cell(f, ix["valor"]).value = VALOR
    ws.cell(f, ix["unidad"]).value = "%"
    ws.cell(f, ix["comentario"]).value = COMENTARIO

    if not escribir:
        print("\n  Vista previa. Para aplicar agrega --escribir\n")
        return

    os.makedirs(os.path.join(AQUI, "respaldos"), exist_ok=True)
    sello = datetime.now().strftime("%Y%m%d_%H%M%S")
    copia = os.path.join(AQUI, "respaldos", "BD_MAESTRA_COCO_antes_proveedor_ventas_pct_junio_" + sello + ".xlsx")
    shutil.copy2(BASE, copia)
    wb.save(BASE)
    print("\n  Respaldo:", os.path.basename(copia))
    print("  Base actualizada.\n")


if __name__ == "__main__":
    main()
