# -*- coding: utf-8 -*-
"""
Corrige la TRM de julio 2026: de 3505.2683 (implicita, autoderivada la semana
pasada) a 3268.95 (la oficial, la misma que ya usa el informe de Burn & Runway
y el modelo financiero V3 -- confirmada porque al reconvertir Colombia con
esta tasa, cada linea del P&G cuadra casi exacto contra el modelo: Admin.
33,017.5 vs 33,017.6; Ventas 36,024.2 vs 36,024.2; Operacion 31,489.1 vs
31,489.1; Costo servicio 98,476.9 vs 98,477.4).

Hace 3 cosas:
  1. Actualiza TRM!2026-07 (trm_cop_usd) de 3505.2683 a 3268.95.
  2. Reconvierte a USD cada pyg_* de Colombia julio (COP / TRM correcta),
     sobrescribiendo el valor que se calculo con la TRM vieja.
  3. Recalcula Consolidado julio (Colombia + EE.UU. + Peru + Costa Rica) para
     cada pyg_* -- las filiales no cambian (ya facturan en USD, no pasan por
     esta TRM), solo cambia el aporte de Colombia.

Uso:  python corregir_trm_julio.py             -> vista previa
      python corregir_trm_julio.py --escribir  -> aplica (con respaldo)
"""
import os
import shutil
import sys
from datetime import datetime

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.join(AQUI, "BD_MAESTRA_COCO.xlsx")
PERIODO = "2026-07"
TRM_VIEJA = 3505.2683
TRM_NUEVA = 3268.95
CODES = ["pyg_servicio_software", "pyg_devoluciones", "pyg_ingresos_operacionales",
          "pyg_costo_ventas", "pyg_utilidad_bruta", "pyg_gasto_administracion",
          "pyg_gasto_proyectos", "pyg_gasto_ventas", "pyg_utilidad_operativa",
          "pyg_ingresos_no_operacionales", "pyg_gastos_financieros", "pyg_utilidad_neta"]


def main():
    escribir = "--escribir" in sys.argv
    wb = openpyxl.load_workbook(BASE)
    ws = wb["BD_Indicadores"]
    H = [str(c.value).strip() if c.value else "" for c in ws[3]]
    ix = {c: H.index(c) + 1 for c in H if c}

    reporte = []

    # --- 1) hoja TRM ---
    wt = wb["TRM"]
    fila_trm = None
    for r in range(1, wt.max_row + 1):
        if str(wt.cell(r, 1).value) == PERIODO:
            fila_trm = r
            break
    if fila_trm:
        antes = wt.cell(fila_trm, 2).value
        reporte.append("  TRM %s: %s -> %s" % (PERIODO, antes, TRM_NUEVA))
        wt.cell(fila_trm, 2).value = TRM_NUEVA
        wt.cell(fila_trm, 3).value = "Oficial julio (confirmada contra modelo V3 y Burn & Runway)"

    # --- 2) reconvertir Colombia julio a USD con la TRM correcta ---
    filas_usd = {}
    for f in range(4, ws.max_row + 1):
        cod = ws.cell(f, ix["codigo_indicador"]).value
        per = ws.cell(f, ix["periodo"]).value
        pais = ws.cell(f, ix["pais"]).value
        seg = ws.cell(f, ix["segmento"]).value
        if cod in CODES and str(per) == PERIODO and pais == "Colombia":
            if seg is None:
                filas_usd.setdefault(cod, {})["cop_fila"] = f
            elif seg == "Consolidación USD":
                filas_usd.setdefault(cod, {})["usd_fila"] = f

    nuevos_colombia = {}
    for cod, d in filas_usd.items():
        if "cop_fila" not in d or "usd_fila" not in d:
            continue
        cop = ws.cell(d["cop_fila"], ix["valor"]).value
        antes = ws.cell(d["usd_fila"], ix["valor"]).value
        nuevo = round(cop / TRM_NUEVA, 2)
        nuevos_colombia[cod] = nuevo
        reporte.append("  Colombia USD %-32s %s -> %s" % (cod, antes, nuevo))
        ws.cell(d["usd_fila"], ix["valor"]).value = nuevo
        ws.cell(d["usd_fila"], ix["comentario"]).value = (
            "Reconvertido con TRM oficial julio (3268.95, antes 3505.2683 autoderivada).")

    # --- 3) recalcular Consolidado julio ---
    existentes = {}
    for f in range(4, ws.max_row + 1):
        cod = ws.cell(f, ix["codigo_indicador"]).value
        per = ws.cell(f, ix["periodo"]).value
        pais = ws.cell(f, ix["pais"]).value
        seg = ws.cell(f, ix["segmento"]).value
        esc = ws.cell(f, ix["escenario"]).value
        if seg == "Consolidación USD" and esc == "Real" and cod and str(per) == PERIODO:
            existentes[(pais, cod)] = f

    siguiente = ws.max_row + 1
    for cod in CODES:
        suma = 0.0
        falta = []
        for pais in ["Colombia", "EE.UU.", "Perú", "Costa Rica"]:
            if pais == "Colombia":
                v = nuevos_colombia.get(cod)
            else:
                f = existentes.get((pais, cod))
                v = ws.cell(f, ix["valor"]).value if f else None
            if v is None:
                falta.append(pais)
            else:
                suma += v
        if falta:
            reporte.append("  AVISO Consolidado %s: falta %s, no se recalcula" % (cod, falta))
            continue
        suma = round(suma, 2)
        clave = ("Consolidado", cod)
        if clave in existentes:
            f = existentes[clave]
            antes = ws.cell(f, ix["valor"]).value
            reporte.append("  Consolidado %-32s %s -> %s" % (cod, antes, suma))
        else:
            f = siguiente
            siguiente += 1
            ws.cell(f, ix["periodo"]).value = PERIODO
            ws.cell(f, ix["pais"]).value = "Consolidado"
            ws.cell(f, ix["compania"]).value = "COCO Tecnologias"
            ws.cell(f, ix["moneda"]).value = "USD"
            ws.cell(f, ix["escenario"]).value = "Real"
            ws.cell(f, ix["segmento"]).value = "Consolidación USD"
            reporte.append("  Consolidado %-32s -> %s (nueva fila)" % (cod, suma))
        ws.cell(f, ix["codigo_indicador"]).value = cod
        ws.cell(f, ix["valor"]).value = suma
        ws.cell(f, ix["unidad"]).value = "USD"
        ws.cell(f, ix["comentario"]).value = "Recalculado: suma Colombia (TRM 3268.95) + EE.UU. + Peru + Costa Rica."

    print("=" * 90)
    print("  CORREGIR TRM JULIO 2026 (3505.2683 -> 3268.95) Y RECALCULAR")
    print("=" * 90)
    print()
    for l in reporte:
        print(l)

    if not escribir:
        print("\n  Vista previa. Para aplicar agrega --escribir\n")
        return

    os.makedirs(os.path.join(AQUI, "respaldos"), exist_ok=True)
    sello = datetime.now().strftime("%Y%m%d_%H%M%S")
    copia = os.path.join(AQUI, "respaldos", "BD_MAESTRA_COCO_antes_trm_julio_" + sello + ".xlsx")
    shutil.copy2(BASE, copia)
    wb.save(BASE)
    print("\n  Respaldo:", os.path.basename(copia))
    print("  Base actualizada.\n")


if __name__ == "__main__":
    main()
