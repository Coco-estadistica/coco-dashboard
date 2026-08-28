# -*- coding: utf-8 -*-
"""
Carga julio 2026 de Gastos Colombia (Administrativos, Proveedores, Comercial,
Operaciones, Desarrollo), calculado por sustraccion:
    julio = (Balance de Prueba ERP, acumulado ene-jul) - (ya conocido ene-jun,
             BD + Anexo2 "Recursos Financieros Informes Julio 2026")
Mapeo de cuenta PUC validado contra el historico ene-jun antes de cargar
(cada categoria calculada cae en el rango normal de esa serie):
    Administrativos (51xx) | Proveedores (72060101 y afines) | Comercial (52xx)
    | Operaciones (72xx, excl. 7206) | Desarrollo (75xx)

Quedan SIN cargar (dan negativo o sin validar => el historico agrupa mas
cuentas de las usadas aqui, no se adivina el resto):
  - Administrativos: Diversos
  - Proveedores: AWS (no aparece ningun tercero con ese nombre en el ERP),
    IVA (sin punto de comparacion historico), Otros (15x el patron normal)
  - Operaciones: Viajes (acumulado ERP < ya cargado ene-jun)

Uso:  python cargar_gastos_colombia_julio.py             -> vista previa
      python cargar_gastos_colombia_julio.py --escribir  -> aplica (con respaldo)
"""
import os
import shutil
import sys
from datetime import datetime

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.join(AQUI, "BD_MAESTRA_COCO.xlsx")
PERIODO = "2026-07"

COMENTARIO = ("Julio 2026 = Balance de Prueba ERP (acumulado ene-jul) menos "
              "ene-jun ya conocido (BD + Anexo2 Recursos Financieros julio 2026). "
              "Cuenta PUC validada contra el rango historico de la serie antes de cargar.")

# (codigo_indicador, detalle, valor)
DATOS = [
    ("gasto_admin", "Personal", 60719000),
    ("gasto_admin", "Honorarios", 7656122),
    ("gasto_admin", "Arrendamiento", 10446087),
    ("gasto_admin", "Seguros", 1206094),
    ("gasto_admin", "Servicios", 3828607),
    ("gasto_admin", "Viajes", 11391185),

    ("costo_proveedor", "Infobip", 93796468),
    ("costo_proveedor", "B2Chat", 20973084),

    ("gasto_comercial", "Personal", 54224968),
    ("gasto_comercial", "Honorario", 12956710),
    ("gasto_comercial", "Seguros", 276332),
    ("gasto_comercial", "Servicios", 13644624),
    ("gasto_comercial", "Viajes", 16944724),
    ("gasto_comercial", "Diversos", 22869473),

    ("costo_operacion", "Personal", 76693508),
    ("costo_operacion", "Honorario", 4710811),
    ("costo_operacion", "Servicios", 10417286),
    ("costo_operacion", "Mantenimientos", 17940044),

    ("gasto_desarrollo", "Personal", 148534887),
    ("gasto_desarrollo", "Honorarios", 4051580),
    ("gasto_desarrollo", "Servicios", 1763472),
    ("gasto_desarrollo", "Viajes", 46625),
]


def main():
    escribir = "--escribir" in sys.argv
    wb = openpyxl.load_workbook(BASE)
    ws = wb["BD_Indicadores"]
    H = [str(c.value).strip() if c.value else "" for c in ws[3]]
    ix = {c: H.index(c) + 1 for c in H if c}

    existentes = {}
    for f in range(4, ws.max_row + 1):
        cod = ws.cell(f, ix["codigo_indicador"]).value
        det = ws.cell(f, ix["detalle"]).value
        per = ws.cell(f, ix["periodo"]).value
        pais = ws.cell(f, ix["pais"]).value
        if cod and pais == "Colombia":
            existentes[(cod, det, str(per))] = f

    siguiente = ws.max_row + 1
    reporte = []
    totales = {}

    for cod, det, valor in DATOS:
        clave = (cod, det, PERIODO)
        totales[cod] = totales.get(cod, 0) + valor
        if clave in existentes:
            f = existentes[clave]
            antes = ws.cell(f, ix["valor"]).value
            reporte.append("  UPDATE %-18s %-14s %s -> %s" % (cod, det, antes, valor))
        else:
            f = siguiente
            siguiente += 1
            ws.cell(f, ix["periodo"]).value = PERIODO
            ws.cell(f, ix["pais"]).value = "Colombia"
            ws.cell(f, ix["compania"]).value = "Coco Colombia"
            ws.cell(f, ix["moneda"]).value = "COP"
            ws.cell(f, ix["escenario"]).value = "Real"
            ws.cell(f, ix["periodicidad"]).value = "Mensual"
            reporte.append("  INSERT %-18s %-14s -> %s" % (cod, det, valor))
        ws.cell(f, ix["codigo_indicador"]).value = cod
        ws.cell(f, ix["detalle"]).value = det
        ws.cell(f, ix["valor"]).value = valor
        ws.cell(f, ix["unidad"]).value = "COP"
        ws.cell(f, ix["comentario"]).value = COMENTARIO

    print("=" * 90)
    print("  CARGA GASTOS COLOMBIA - JULIO 2026 (Administrativos/Proveedores/Comercial/Operaciones/Desarrollo)")
    print("=" * 90)
    print()
    for l in reporte:
        print(l)
    print()
    for cod, tot in totales.items():
        print("  Total %-18s julio: %s" % (cod, format(tot, ",")))
    print("\n  SIN cargar (ver docstring): gasto_admin/Diversos, costo_proveedor/AWS-IVA-Otros, "
          "costo_operacion/Viajes\n")

    if not escribir:
        print("  Vista previa. Para aplicar agrega --escribir\n")
        return

    os.makedirs(os.path.join(AQUI, "respaldos"), exist_ok=True)
    sello = datetime.now().strftime("%Y%m%d_%H%M%S")
    copia = os.path.join(AQUI, "respaldos", "BD_MAESTRA_COCO_antes_gastos_colombia_julio_" + sello + ".xlsx")
    shutil.copy2(BASE, copia)
    wb.save(BASE)
    print("  Respaldo:", os.path.basename(copia))
    print("  Base actualizada.\n")


if __name__ == "__main__":
    main()
