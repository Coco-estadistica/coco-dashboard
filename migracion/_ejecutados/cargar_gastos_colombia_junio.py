# -*- coding: utf-8 -*-
"""
Carga junio 2026 de Gastos Colombia (Administrativos/Diversos, Proveedores,
Comercial, Operaciones, Desarrollo) -- dato DIRECTO de la tabla dinamica del
Anexo 2 "Recursos Financieros Informes Julio 2026" (hoja Noviembre, ultima
fila de cada bloque = Junio). A diferencia de julio, junio no se calcula por
resta: viene tal cual del archivo, asi que se carga completo.

Fuente por bloque (hoja "Noviembre"):
  Administrativos fila 296 (Diversos, unico dato que faltaba -- el resto de
    junio ya estaba cargado)
  Proveedores      fila 330
  Comercial        fila 435
  Operaciones      fila 366
  Desarrollo       fila 400

Uso:  python cargar_gastos_colombia_junio.py             -> vista previa
      python cargar_gastos_colombia_junio.py --escribir  -> aplica (con respaldo)
"""
import os
import shutil
import sys
from datetime import datetime

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.join(AQUI, "BD_MAESTRA_COCO.xlsx")
PERIODO = "2026-06"

COMENTARIO = ("Junio 2026, dato directo de Anexo 2 Recursos Financieros Informes "
              "Julio 2026 (hoja Noviembre, fila de Junio de cada bloque). No calculado.")

# (codigo_indicador, detalle, valor)
DATOS = [
    ("gasto_admin", "Diversos", 21100000),

    ("costo_proveedor", "Infobip", 110800000),
    ("costo_proveedor", "IVA", 0),
    ("costo_proveedor", "AWS", 40900000),
    ("costo_proveedor", "B2Chat", 19800000),
    ("costo_proveedor", "Otros", 16800000),

    ("gasto_comercial", "Personal", 54200000),
    ("gasto_comercial", "Honorario", 43200000),
    ("gasto_comercial", "Seguros", 800000),
    ("gasto_comercial", "Servicios", 22400000),
    ("gasto_comercial", "Viajes", 4400000),
    ("gasto_comercial", "Diversos", 22300000),

    ("costo_operacion", "Personal", 81400000),
    ("costo_operacion", "Honorario", 0),
    ("costo_operacion", "Servicios", 12300000),
    ("costo_operacion", "Mantenimientos", 7000000),
    ("costo_operacion", "Viajes", 8100000),

    ("gasto_desarrollo", "Personal", 153900000),
    ("gasto_desarrollo", "Honorarios", 3300000),
    ("gasto_desarrollo", "Servicios", 2600000),
    ("gasto_desarrollo", "Viajes", 0),
]


def main():
    escribir = "--escribir" in sys.argv
    wb = openpyxl.load_workbook(BASE)
    ws = wb["BD_Indicadores"]
    H = [str(c.value).strip() if c.value else "" for c in ws[3]]
    ix = {c: H.index(c) + 1 for c in H if c}

    existentes = {}
    for f in range(4, ws.max_row + 1):
        cod = ws.cell(f, ix["codigo_indicador"]).value
        det = ws.cell(f, ix["detalle"]).value
        per = ws.cell(f, ix["periodo"]).value
        pais = ws.cell(f, ix["pais"]).value
        if cod and pais == "Colombia":
            existentes[(cod, det, str(per))] = f

    siguiente = ws.max_row + 1
    reporte = []
    totales = {}

    for cod, det, valor in DATOS:
        clave = (cod, det, PERIODO)
        totales[cod] = totales.get(cod, 0) + valor
        if clave in existentes:
            f = existentes[clave]
            antes = ws.cell(f, ix["valor"]).value
            reporte.append("  UPDATE %-18s %-14s %s -> %s" % (cod, det, antes, valor))
        else:
            f = siguiente
            siguiente += 1
            ws.cell(f, ix["periodo"]).value = PERIODO
            ws.cell(f, ix["pais"]).value = "Colombia"
            ws.cell(f, ix["compania"]).value = "Coco Colombia"
            ws.cell(f, ix["moneda"]).value = "COP"
            ws.cell(f, ix["escenario"]).value = "Real"
            ws.cell(f, ix["periodicidad"]).value = "Mensual"
            reporte.append("  INSERT %-18s %-14s -> %s" % (cod, det, valor))
        ws.cell(f, ix["codigo_indicador"]).value = cod
        ws.cell(f, ix["detalle"]).value = det
        ws.cell(f, ix["valor"]).value = valor
        ws.cell(f, ix["unidad"]).value = "COP"
        ws.cell(f, ix["comentario"]).value = COMENTARIO

    print("=" * 90)
    print("  CARGA GASTOS COLOMBIA - JUNIO 2026 (completar huecos, dato directo de Anexo2)")
    print("=" * 90)
    print()
    for l in reporte:
        print(l)
    print()
    for cod, tot in totales.items():
        print("  Total %-18s junio (solo lo cargado aqui): %s" % (cod, format(tot, ",")))

    if not escribir:
        print("\n  Vista previa. Para aplicar agrega --escribir\n")
        return

    os.makedirs(os.path.join(AQUI, "respaldos"), exist_ok=True)
    sello = datetime.now().strftime("%Y%m%d_%H%M%S")
    copia = os.path.join(AQUI, "respaldos", "BD_MAESTRA_COCO_antes_gastos_colombia_junio_" + sello + ".xlsx")
    shutil.copy2(BASE, copia)
    wb.save(BASE)
    print("\n  Respaldo:", os.path.basename(copia))
    print("  Base actualizada.\n")


if __name__ == "__main__":
    main()
