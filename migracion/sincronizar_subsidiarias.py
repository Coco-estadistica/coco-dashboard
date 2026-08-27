# -*- coding: utf-8 -*-
"""
Sincroniza Subsidiarias_PyG con BD_Indicadores (la hoja que lee el dashboard).

Motivo
------
`conciliar_capas.py` detectó que Subsidiarias_PyG conservaba la cifra anterior de
Costa Rica junio (ingresos 137.010 / utilidad 134.338), reversada en su momento
porque el auxiliar de la filial no registraba ingresos ese mes. La corrección se
aplicó a BD_Indicadores y BD_PYG_OFICIAL pero no llegó a esta capa.

Qué hace
--------
Recorre Subsidiarias_PyG y, para cada fila, toma el valor vigente de BD_Indicadores
(bloque 'Consolidación USD', misma moneda). Solo escribe donde hay diferencia.
No inventa datos: si BD_Indicadores no tiene la cifra, deja la fila intacta y la reporta.

ANTES DE USARLO, LEER ESTO
--------------------------
El encabezado de conciliar_capas.py dice que Subsidiarias_PyG y BD_PYG_OFICIAL son
archivo historico y quedaron desalineadas A PROPOSITO el 13-ago-2026, cuando se
reemplazo el P&G de EE.UU./Costa Rica/Consolidado por cifras homologadas desde los
auxiliares contables. Las decenas de diferencias de los checks B y C son esperadas.
El check que importa es el A (Paises vs Consolidado) y debe seguir en OK.

Es decir: correr esto BORRA el antes de esa homologacion. Se ejecuto por error el
27-ago-2026 y se revirtio desde el respaldo. Usarlo solo si alguien decide, a
sabiendas, que esa capa debe reflejar las cifras vigentes.

Uso:  python sincronizar_subsidiarias.py            -> vista previa
      python sincronizar_subsidiarias.py --escribir -> aplica (con respaldo)
"""
import os
import shutil
import sys
from datetime import datetime

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
# Las dos capas viven en libros distintos desde que la base se separo el 13-ago-2026:
# BD_Indicadores (la que manda) en la base maestra, y Subsidiarias_PyG en el libro de
# trazabilidad. Antes estaban en el mismo archivo y este script abria uno solo; desde
# la separacion moria con KeyError y Corregir_Base.bat dejo de funcionar sin que nadie
# lo notara, porque el .bat no distingue el error de una corrida limpia.
BASE = os.path.join(AQUI, "BD_MAESTRA_COCO.xlsx")
TRAZA = os.path.join(AQUI, "BD_TRAZABILIDAD_COCO.xlsx")
SEG = "Consolidación USD"
TOL = 1.0


def encabezado(ws):
    for r in range(1, 8):
        fila = [str(c.value).strip().lower() if c.value else "" for c in ws[r]]
        if "periodo" in fila and "codigo_indicador" in fila:
            return r
    raise SystemExit(f"No encuentro el encabezado en {ws.title}")


def main():
    escribir = "--escribir" in sys.argv
    wb = openpyxl.load_workbook(BASE)

    # --- índice de BD_Indicadores (fuente de verdad) ---
    wsb = wb["BD_Indicadores"]
    hb = encabezado(wsb)
    Hb = [str(c.value).strip() if c.value else "" for c in wsb[hb]]
    ib = {c: Hb.index(c) for c in Hb if c}
    vigente = {}
    for fila in wsb.iter_rows(min_row=hb + 1, values_only=True):
        if not fila or not fila[ib["codigo_indicador"]]:
            continue
        if (fila[ib["segmento"]] or "") != SEG:
            continue
        clave = (str(fila[ib["periodo"]]), fila[ib["pais"]], fila[ib["codigo_indicador"]])
        vigente[clave] = fila[ib["valor"]]

    # --- recorrer Subsidiarias_PyG ---
    # Subsidiarias_PyG vive en el libro de trazabilidad, no en la base maestra.
    wbt = openpyxl.load_workbook(TRAZA, data_only=False)
    if "Subsidiarias_PyG" not in wbt.sheetnames:
        print("No encuentro Subsidiarias_PyG en %s" % os.path.basename(TRAZA))
        return
    wss = wbt["Subsidiarias_PyG"]
    hs = encabezado(wss)
    Hs = [str(c.value).strip() if c.value else "" for c in wss[hs]]
    isub = {c: Hs.index(c) for c in Hs if c}
    col_valor = isub["valor"] + 1

    cambios, sin_par = [], []
    for r in range(hs + 1, wss.max_row + 1):
        cod = wss.cell(row=r, column=isub["codigo_indicador"] + 1).value
        if not cod:
            continue
        per = str(wss.cell(row=r, column=isub["periodo"] + 1).value)
        pais = wss.cell(row=r, column=isub["pais"] + 1).value
        actual = wss.cell(row=r, column=col_valor).value
        nuevo = vigente.get((per, pais, cod))
        if nuevo is None:
            sin_par.append((per, pais, cod))
            continue
        if isinstance(actual, (int, float)) and isinstance(nuevo, (int, float)) and abs(actual - nuevo) > TOL:
            cambios.append((r, per, pais, cod, actual, nuevo))
            if escribir:
                wss.cell(row=r, column=col_valor, value=nuevo)

    print("=" * 84)
    print("  SINCRONIZAR Subsidiarias_PyG <- BD_Indicadores —",
          "ESCRITURA" if escribir else "VISTA PREVIA (no escribe)")
    print("=" * 84)
    if cambios:
        print(f"  {'fila':>5}  {'periodo':<9}{'país':<13}{'indicador':<30}{'antes':>15}{'después':>15}")
        print("  " + "-" * 80)
        for r, per, pais, cod, a, n in cambios:
            print(f"  {r:>5}  {per:<9}{str(pais):<13}{cod:<30}{a:>15,.2f}{n:>15,.2f}")
    else:
        print("  Sin diferencias: las dos capas ya están alineadas.")
    if sin_par:
        print(f"\n  {len(sin_par)} fila(s) sin equivalente en BD_Indicadores (se dejan intactas):")
        for per, pais, cod in sin_par[:8]:
            print(f"     {per}  {pais}  {cod}")
        if len(sin_par) > 8:
            print(f"     … y {len(sin_par)-8} más")

    if not escribir:
        print("\n  Nada se escribió. Para aplicar:  python sincronizar_subsidiarias.py --escribir\n")
        return
    if not cambios:
        return

    os.makedirs(os.path.join(AQUI, "respaldos"), exist_ok=True)
    sello = datetime.now().strftime("%Y%m%d_%H%M%S")
    # Se respalda y se escribe el libro de TRAZABILIDAD: la base maestra no se toca,
    # porque BD_Indicadores es la fuente y aqui solo se alinea la capa de archivo.
    copia = os.path.join(AQUI, "respaldos", f"BD_TRAZABILIDAD_COCO_antes_sync_subsidiarias_{sello}.xlsx")
    shutil.copy2(TRAZA, copia)
    wbt.save(TRAZA)
    print(f"\n  Respaldo: {os.path.basename(copia)}")
    print(f"  Aplicados {len(cambios)} cambio(s). Vuelve a correr conciliar_capas.py para confirmar.\n")


if __name__ == "__main__":
    main()
