# -*- coding: utf-8 -*-
"""Carga Plantilla_Mensual_MRR_Cartera.xlsx en BD_MAESTRA_COCO.xlsx.

PARA QUE
--------
El cierre de MRR y de cartera se venia haciendo con un script distinto cada mes
(cargar_mrr_julio.py, cargar_cartera_julio.py, cargar_mrr_cartera_agosto.py...).
Cada uno con sus cifras adentro. Esto lo reemplaza por: se diligencia la
plantilla, se corre esto, y el script no escribe si algo no cuadra.

USO
---
    python migracion/cargar_plantilla_mensual.py <plantilla.xlsx>              (en seco)
    python migracion/cargar_plantilla_mensual.py <plantilla.xlsx> --escribir   (aplica)

En seco muestra fila por fila que pasaria --alta, cambio o sin cambio-- y no toca
la base. Con --escribir respalda la base y aplica.

LOS CHEQUEOS
------------
Se vuelven a calcular AQUI, en Python, sobre las cifras crudas. No se leen los de
la hoja CHEQUEOS: openpyxl lee el ultimo valor que Excel dejo guardado, y si el
archivo se diligencio y se mando sin abrirlo, esos valores estan viejos. Un
chequeo que se cree cumplido sin serlo es peor que no tenerlo.
"""
import argparse
import os
import shutil
import sys

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.join(AQUI, "BD_MAESTRA_COCO.xlsx")

COLS = ["periodo", "pais", "compania", "moneda", "escenario", "codigo_indicador",
        "segmento", "detalle", "valor", "unidad", "comentario"]

# Codigos que llevan varias filas por periodo: la llave incluye 'detalle'.
POR_DETALLE = {"cartera_vencida_cliente", "distribucion_ingresos"}
# Codigos que llevan una fila por pais.
POR_PAIS = {"liquidez_total"}

TOL = 1.0


def leer_plantilla(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    if "BD_Indicadores" not in wb.sheetnames:
        raise SystemExit("\nLa plantilla no tiene la hoja BD_Indicadores.")
    ws = wb["BD_Indicadores"]
    cab = [str(c.value or "").strip() for c in ws[1]]
    ix = {c: i for i, c in enumerate(cab) if c}
    faltan = [c for c in COLS if c not in ix]
    if faltan:
        raise SystemExit("\nA la hoja BD_Indicadores le faltan columnas: %s" % ", ".join(faltan))

    filas = []
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, i + 1).value for i in range(len(cab))]
        v = vals[ix["valor"]]
        if v is None or str(v).strip() == "":
            continue                      # fila pre-poblada sin diligenciar: se ignora
        f = {c: (vals[ix[c]] if ix[c] < len(vals) else None) for c in COLS}
        f["periodo"] = str(f["periodo"] or "").strip()
        for c in ("pais", "compania", "moneda", "escenario", "codigo_indicador",
                  "segmento", "detalle", "unidad", "comentario"):
            f[c] = str(f[c] or "").strip()
        try:
            f["valor"] = float(f["valor"])
        except (TypeError, ValueError):
            raise SystemExit("\nFila %d: el valor «%s» no es un numero." % (r, f["valor"]))
        f["_fila"] = r
        filas.append(f)
    return filas


def sumar(filas, codigo):
    t, hubo = 0.0, False
    for f in filas:
        if f["codigo_indicador"] == codigo:
            t += f["valor"]
            hubo = True
    return t if hubo else None


def chequear(filas, mrr_anterior):
    """Devuelve [(nivel, texto)]. nivel ALTO frena la escritura; AVISO no."""
    out = []
    if not filas:
        return [("ALTO", "La plantilla no trae ninguna fila diligenciada.")]

    pers = sorted({f["periodo"] for f in filas})
    if len(pers) != 1:
        out.append(("ALTO", "Hay %d periodos distintos en el archivo (%s). Una carga, un mes."
                    % (len(pers), ", ".join(pers))))
    for f in filas:
        if not f["periodo"] or len(f["periodo"]) != 7 or f["periodo"][4] != "-":
            out.append(("ALTO", "Fila %d: periodo «%s» no tiene formato AAAA-MM."
                        % (f["_fila"], f["periodo"])))
            break

    # duplicados de llave
    vistas = {}
    for f in filas:
        k = (f["periodo"], f["codigo_indicador"], f["pais"],
             f["detalle"] if f["codigo_indicador"] in POR_DETALLE else "")
        if k in vistas:
            out.append(("ALTO", "Fila %d repite la llave de la fila %d (%s / %s / %s). "
                                "La segunda pisaria a la primera."
                        % (f["_fila"], vistas[k], f["codigo_indicador"], f["pais"], f["detalle"] or "-")))
        vistas[k] = f["_fila"]

    # filas que deben llevar detalle
    for f in filas:
        if f["codigo_indicador"] in POR_DETALLE and not f["detalle"]:
            out.append(("ALTO", "Fila %d: «%s» necesita la columna 'detalle' diligenciada."
                        % (f["_fila"], f["codigo_indicador"])))

    # signos
    for cod in ("contraction_mrr", "churned_mrr"):
        v = sumar(filas, cod)
        if v is not None and v > 0:
            out.append(("ALTO", "«%s» viene en positivo (%.0f). Debe ir negativo o el "
                                "puente suma donde debia restar." % (cod, v)))

    # puente de MRR
    mrr = sumar(filas, "mrr_final")
    movs = [sumar(filas, c) for c in ("new_mrr", "reactivacion_mrr", "expansion_mrr",
                                      "contraction_mrr", "churned_mrr")]
    if mrr is not None and any(m is not None for m in movs):
        if mrr_anterior is None:
            out.append(("AVISO", "No se pudo comprobar el puente de MRR: no encontre el "
                                 "mrr_final del mes anterior en la base. Pasalo con --mrr-anterior."))
        else:
            esperado = mrr_anterior + sum(m or 0 for m in movs)
            if abs(esperado - mrr) > TOL:
                out.append(("ALTO", "El puente de MRR no cierra: %,.0f (anterior + movimientos) "
                                    "contra %,.0f (mrr_final). Diferencia %,.0f."
                            .replace(",", "") % (esperado, mrr, esperado - mrr)))

    # cartera concilia
    tot, sana, ven = sumar(filas, "cartera_total"), sumar(filas, "cartera_sana"), sumar(filas, "cartera_vencida")
    if None not in (tot, sana, ven) and abs((sana + ven) - tot) > TOL:
        out.append(("ALTO", "Cartera no concilia: sana + vencida = %.0f contra cartera_total = %.0f. "
                            "La morosidad saldria distinta segun cual se use." % (sana + ven, tot)))

    # el desglose por cliente no puede pasarse de la vencida
    porcli = sumar(filas, "cartera_vencida_cliente")
    if porcli is not None and ven is not None and porcli > ven * 1.01:
        out.append(("ALTO", "El desglose por cliente suma %.0f y la cartera vencida es %.0f."
                    % (porcli, ven)))

    # moneda coherente con el pais
    esperada = {"Colombia": "COP", "EE.UU.": "USD", "Perú": "PEN", "Costa Rica": "USD"}
    for f in filas:
        e = esperada.get(f["pais"])
        if e and f["moneda"] and f["moneda"] != e and f["unidad"] not in ("%", "días"):
            out.append(("AVISO", "Fila %d: %s en %s (lo habitual es %s). Si ya la convertiste, "
                                 "el tablero la volveria a convertir."
                        % (f["_fila"], f["pais"], f["moneda"], e)))

    # ARR
    arr = sumar(filas, "arr_mrr")
    if arr is not None and mrr:
        if abs(arr - mrr * 12) / (mrr * 12) > 0.15:
            out.append(("AVISO", "arr_mrr (%.0f) se aparta mas de 15%% de mrr_final x 12 (%.0f)."
                        % (arr, mrr * 12)))
    return out


def leer_base():
    if not os.path.exists(BASE):
        raise SystemExit("\nNo encuentro %s" % BASE)
    wb = openpyxl.load_workbook(BASE)
    ws = wb["BD_Indicadores"]
    H = [str(c.value).strip() if c.value else "" for c in ws[3]]
    ix = {c: H.index(c) + 1 for c in H if c}
    return wb, ws, ix


def mrr_anterior_de_base(ws, ix, periodo):
    """mrr_final del mes inmediatamente anterior, para comprobar el puente."""
    a, m = int(periodo[:4]), int(periodo[5:7])
    m -= 1
    if m == 0:
        a, m = a - 1, 12
    prev = "%04d-%02d" % (a, m)
    for r in range(4, ws.max_row + 1):
        if (str(ws.cell(r, ix["codigo_indicador"]).value or "").strip() == "mrr_final"
                and str(ws.cell(r, ix["periodo"]).value) == prev
                and str(ws.cell(r, ix["pais"]).value or "").strip() == "Colombia"):
            v = ws.cell(r, ix["valor"]).value
            if v is not None:
                return float(v), prev
    return None, prev


def buscar_fila(ws, ix, f):
    por_det = f["codigo_indicador"] in POR_DETALLE
    for r in range(4, ws.max_row + 1):
        if str(ws.cell(r, ix["codigo_indicador"]).value or "").strip() != f["codigo_indicador"]:
            continue
        if str(ws.cell(r, ix["periodo"]).value) != f["periodo"]:
            continue
        if str(ws.cell(r, ix["pais"]).value or "").strip() != f["pais"]:
            continue
        if por_det and str(ws.cell(r, ix["detalle"]).value or "").strip() != f["detalle"]:
            continue
        return r
    return None


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("plantilla")
    ap.add_argument("--escribir", action="store_true", help="aplica a la base (con respaldo)")
    ap.add_argument("--mrr-anterior", type=float, default=None,
                    help="mrr_final del mes anterior, si la base no lo tiene")
    args = ap.parse_args()

    if not os.path.exists(args.plantilla):
        print("No encuentro %s" % args.plantilla)
        return 2

    filas = leer_plantilla(args.plantilla)
    print("\n  %s: %d fila(s) diligenciada(s)" % (os.path.basename(args.plantilla), len(filas)))
    if not filas:
        print("  No hay nada que cargar.")
        return 1

    wb, ws, ix = leer_base()
    per = filas[0]["periodo"]
    mrr_ant, mes_prev = (args.mrr_anterior, "dado a mano") if args.mrr_anterior is not None \
        else mrr_anterior_de_base(ws, ix, per)
    if mrr_ant is not None:
        print("  MRR del mes anterior (%s): %s" % (mes_prev, "{:,.0f}".format(mrr_ant)))

    hallazgos = chequear(filas, mrr_ant)
    altos = [h for h in hallazgos if h[0] == "ALTO"]
    if hallazgos:
        print("\n  CHEQUEOS")
        for niv, txt in hallazgos:
            print("    [%s] %s" % (niv, txt))
    else:
        print("\n  CHEQUEOS: todo en orden.")

    print("\n  PROPUESTA")
    print("  %-26s %-12s %-16s %18s %18s  %s"
          % ("codigo", "pais", "detalle", "en la base", "en la plantilla", "que pasaria"))
    print("  " + "-" * 116)
    altas = cambios = igual = 0
    for f in sorted(filas, key=lambda x: (x["codigo_indicador"], x["pais"], x["detalle"])):
        r = buscar_fila(ws, ix, f)
        antes = ws.cell(r, ix["valor"]).value if r else None
        antes = float(antes) if isinstance(antes, (int, float)) else None
        if r is None:
            que, altas = "alta", altas + 1
        elif antes is None or abs(antes - f["valor"]) > 0.005:
            que, cambios = "CAMBIA", cambios + 1
        else:
            que, igual = "sin cambio", igual + 1
        fmt = lambda v: "—" if v is None else "{:,.2f}".format(v)
        print("  %-26s %-12s %-16s %18s %18s  %s"
              % (f["codigo_indicador"][:26], f["pais"][:12], (f["detalle"] or "")[:16],
                 fmt(antes), fmt(f["valor"]), que))
    print("  " + "-" * 116)
    print("  %d alta(s) · %d cambio(s) · %d sin cambio" % (altas, cambios, igual))

    if altos:
        print("\n  HAY %d CONTROL(ES) EN ALTO. No se escribe nada hasta resolverlos." % len(altos))
        return 1
    if not args.escribir:
        print("\n  Nada se ha escrito. Aplica con:  python migracion/cargar_plantilla_mensual.py %s --escribir"
              % os.path.basename(args.plantilla))
        return 0

    resp = BASE.replace(".xlsx", "_antes_%s.xlsx" % per.replace("-", ""))
    shutil.copy2(BASE, resp)
    n = 0
    for f in filas:
        r = buscar_fila(ws, ix, f)
        if r is None:
            r = ws.max_row + 1
            for col in ("periodo", "pais", "compania", "moneda", "escenario",
                        "codigo_indicador", "segmento", "detalle", "unidad"):
                if f[col] != "":
                    ws.cell(r, ix[col]).value = f[col]
        ws.cell(r, ix["valor"]).value = f["valor"]
        if f["comentario"]:
            ws.cell(r, ix["comentario"]).value = f["comentario"]
        n += 1
    wb.save(BASE)
    print("\n  Escrito: %d fila(s). Respaldo: %s" % (n, os.path.basename(resp)))
    print("  Recarga el tablero con Ctrl+F5 (o «Cargar BD») para verlo.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
