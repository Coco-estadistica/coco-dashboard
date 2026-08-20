# -*- coding: utf-8 -*-
"""
Corrige la TRM de julio 2026 en la CAPA DE ARCHIVO (BD_PYG_OFICIAL, libro de
trazabilidad). Es el hermano de corregir_trm_julio.py, que ya hizo lo mismo en
BD_MAESTRA_COCO.xlsx el 14-ago-2026.

Por que existe
--------------
Cuando se corrigio la TRM de julio de 3505.2683 (autoderivada) a 3268.95 (la
oficial, confirmada contra el modelo V3 y Burn & Runway), se reconvirtio
BD_Indicadores pero NO BD_PYG_OFICIAL. La capa de archivo quedo con el julio de
Colombia convertido a la tasa de JUNIO.

La huella es inconfundible: las 12 lineas del P&G de Colombia julio difieren
contra BD_Indicadores con un ratio identico de 1.072292 = 3505.2683 / 3268.95.
Un ratio constante en todos los rubros es diferencia de TASA, no de cifra.

Esto NO es el desalineamiento por diseno del 13-ago-2026 (EE.UU./Costa Rica/Peru
con P&G homologado desde auxiliares, que se dejo congelado a proposito). Es un
error de conversion que quedo suelto, y se corrige.

Alcance -- deliberadamente estrecho
-----------------------------------
  1. Colombia 2026-07, segmento "Consolidacion USD": se RECALCULA desde la fila
     COP de la misma hoja (COP / 3268.95). No se copia de BD_Indicadores: la
     regla del proyecto es que Colombia se origina en COP. El resultado coincide
     al centavo con BD_Indicadores -- eso es la verificacion, no el metodo.
  2. Consolidado 2026-07: se ajusta SUMANDO el delta de Colombia, no igualandolo
     a BD_Indicadores. Igualarlo tambien absorberia las diferencias de
     EE.UU./Peru que estan congeladas por diseno, y borraria el rastro de
     auditoria que esa capa existe para conservar.

Las filas COP no se tocan: son el dato oficial del PDF y estan correctas.
Las filiales no se tocan: facturan en USD y no pasan por esta TRM.

Uso:  python corregir_trm_julio_trazabilidad.py             -> vista previa
      python corregir_trm_julio_trazabilidad.py --escribir  -> aplica (con respaldo)
"""
import os
import shutil
import sys
from datetime import datetime

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
TRAZA = os.path.join(AQUI, "BD_TRAZABILIDAD_COCO.xlsx")
RESPALDOS = os.path.join(AQUI, "respaldos")
HOJA = "BD_PYG_OFICIAL"
PERIODO = "2026-07"
SEG = "Consolidación USD"
TRM_VIEJA = 3505.2683
TRM_NUEVA = 3268.95
CODES = ["pyg_servicio_software", "pyg_devoluciones", "pyg_ingresos_operacionales",
         "pyg_costo_ventas", "pyg_utilidad_bruta", "pyg_gasto_administracion",
         "pyg_gasto_proyectos", "pyg_gasto_ventas", "pyg_utilidad_operativa",
         "pyg_ingresos_no_operacionales", "pyg_gastos_financieros", "pyg_utilidad_neta"]
NOTA = (f"Reconvertido con TRM oficial julio ({TRM_NUEVA}, antes {TRM_VIEJA} "
        f"autoderivada). Corregido {datetime.now():%Y-%m-%d} en la capa de archivo; "
        f"BD_Indicadores ya estaba correcta.")


def main():
    escribir = "--escribir" in sys.argv
    if not os.path.exists(TRAZA):
        sys.exit(f"No encuentro el libro de trazabilidad: {TRAZA}")

    wb = openpyxl.load_workbook(TRAZA)
    ws = wb[HOJA]
    H = [str(c.value).strip() if c.value else "" for c in ws[1]]
    ix = {c: H.index(c) + 1 for c in H if c}

    # ---- indexar las filas de julio ----
    cop, usd = {}, {}          # cop[codigo] = valor COP ; usd[(pais, codigo)] = fila
    for r in range(2, ws.max_row + 1):
        per = ws.cell(r, ix["periodo"]).value
        if str(per) != PERIODO:
            continue
        pais = ws.cell(r, ix["pais"]).value
        cod = ws.cell(r, ix["codigo_indicador"]).value
        seg = ws.cell(r, ix["segmento"]).value or ""
        if cod not in CODES:
            continue
        if pais == "Colombia" and seg == "":
            cop[cod] = ws.cell(r, ix["valor"]).value
        elif seg == SEG and pais in ("Colombia", "Consolidado"):
            usd[(pais, cod)] = r

    faltan = [c for c in CODES if c not in cop]
    if faltan:
        sys.exit(f"Faltan filas COP de Colombia {PERIODO}: {faltan}. No se corrige a ciegas.")

    # ---- calcular ----
    cambios = []
    for cod in CODES:
        fc, fk = usd.get(("Colombia", cod)), usd.get(("Consolidado", cod))
        if fc is None or fk is None:
            sys.exit(f"Falta fila USD de {cod} en {PERIODO}. No se corrige a ciegas.")
        viejo_col = ws.cell(fc, ix["valor"]).value
        nuevo_col = round(cop[cod] / TRM_NUEVA, 2)
        delta = round(nuevo_col - viejo_col, 2)
        viejo_con = ws.cell(fk, ix["valor"]).value
        nuevo_con = round(viejo_con + delta, 2)
        cambios.append((cod, fc, viejo_col, nuevo_col, fk, viejo_con, nuevo_con, delta))

    print("=" * 92)
    print(f"  TRM JULIO EN LA CAPA DE ARCHIVO — {HOJA} ({os.path.basename(TRAZA)})")
    print(f"  {TRM_VIEJA} (junio)  ->  {TRM_NUEVA} (oficial julio)")
    print("=" * 92)
    print(f"  {'codigo':32s} {'Colombia antes':>15} {'Colombia nuevo':>15} {'delta':>12}")
    for cod, _, va, vn, _, ka, kn, d in cambios:
        print(f"  {cod:32s} {va:15,.2f} {vn:15,.2f} {d:12,.2f}")
    print(f"\n  {'codigo':32s} {'Consol. antes':>15} {'Consol. nuevo':>15}")
    for cod, _, _, _, _, ka, kn, _ in cambios:
        print(f"  {cod:32s} {ka:15,.2f} {kn:15,.2f}")

    if not escribir:
        print("\n  VISTA PREVIA — no se escribio nada.")
        print("  Para aplicar:  python corregir_trm_julio_trazabilidad.py --escribir")
        return

    # ---- respaldo antes de tocar nada (regla 6 del proyecto) ----
    os.makedirs(RESPALDOS, exist_ok=True)
    sello = datetime.now().strftime("%Y%m%d_%H%M%S")
    destino = os.path.join(RESPALDOS, f"BD_TRAZABILIDAD_COCO_antes_trm_julio_oficial_{sello}.xlsx")
    shutil.copy2(TRAZA, destino)
    print(f"\n  Respaldo: respaldos/{os.path.basename(destino)}")

    for cod, fc, _, vn, fk, _, kn, _ in cambios:
        ws.cell(fc, ix["valor"]).value = vn
        ws.cell(fk, ix["valor"]).value = kn
        for fila in (fc, fk):
            prev = ws.cell(fila, ix["comentario"]).value or ""
            prev = prev.replace(str(TRM_VIEJA), str(TRM_NUEVA))
            ws.cell(fila, ix["comentario"]).value = f"{prev} {NOTA}".strip()

    wb.save(TRAZA)
    print(f"  Escrito: {len(cambios) * 2} celdas de valor + comentarios en {HOJA}.")
    print("  Verifica con:  python conciliar_capas.py")


if __name__ == "__main__":
    main()
