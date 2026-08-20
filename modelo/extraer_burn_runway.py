# -*- coding: utf-8 -*-
"""
Extrae la hoja "Burn & RunW" del Modelo y genera burn_runway.json para el dashboard.

Por qué un JSON y no filas en la base:
  La base (BD_Indicadores) guarda HECHOS (un indicador, un periodo, un valor). Este
  informe es narrativo: tablas con lecturas, escenarios y proyecciones. Meterlo como
  hechos lo desarmaría. Se publica como un archivo aparte que el dashboard lee, y se
  regenera cada vez que actualices el modelo.

Uso:
    python extraer_burn_runway.py
    (luego abre el dashboard: la pestaña "Burn & Runway" se actualiza sola)
"""
import hashlib
import json
import os
import re
import sys
from datetime import datetime

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
MODELO = os.path.join(AQUI, "Modelo_COCO.xlsx")
SALIDA = os.path.join(AQUI, "..", "burn_runway.json")
HOJA = "Burn & RunW"

# Ensayo contra un modelo candidato sin reemplazar nada:
#   python extraer_burn_runway.py "ruta\\del\\modelo.xlsx" prueba.json
if len(sys.argv) > 1:
    MODELO = sys.argv[1]
if len(sys.argv) > 2:
    SALIDA = sys.argv[2]


def huella(ruta):
    """Huella del modelo del que salió este informe.

    Sin esto no hay forma de saber que el modelo cambió: la pestaña seguiría mostrando
    las cifras del corte viejo indefinidamente y con toda la apariencia de estar al día.
    Es el mismo problema que teníamos con las capas duplicadas de la base.
    """
    h = hashlib.md5()
    with open(ruta, "rb") as f:
        for trozo in iter(lambda: f.read(1 << 20), b""):
            h.update(trozo)
    st = os.stat(ruta)
    return {
        "archivo": os.path.basename(ruta),
        "bytes": st.st_size,
        "modificado": datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M"),
        "md5": h.hexdigest(),
    }


def main():
    wb = openpyxl.load_workbook(MODELO, data_only=True)
    ws = wb[HOJA]

    MES = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]

    def c(col, fila):
        v = ws[f"{col}{fila}"].value
        if isinstance(v, str):
            v = v.strip()
            if v.startswith("#"):      # errores de fórmula (#NAME?, #REF!) -> se ignoran
                return None
        elif isinstance(v, datetime):
            # varias etiquetas de mes quedaron como fecha en el modelo -> "Abr-26"
            return f"{MES[v.month - 1].capitalize()}-{str(v.year)[2:]}"
        return v

    # ---------- anclaje por etiquetas, no por número de fila ----------
    # El modelo gana una fila cada vez que se cierra un mes. Con posiciones fijas, el
    # informe no se rompe: se desplaza. Lee la celda de al lado y produce cifras
    # plausibles pero equivocadas, que es el peor resultado posible.
    LIM = 220

    def norm(v):
        return " ".join(str(v).split()).lower() if v is not None else ""

    def fila_con(texto, col="B", desde=1, exacto=False, ultima=False):
        """`ultima`: la hoja repite algunos títulos (un resumen corto arriba y el
        completo abajo). Los bloques buenos son los de abajo, que es adonde apuntaban
        las filas fijas del código original."""
        t = texto.lower()
        hallada = None
        for f in range(desde, LIM + 1):
            e = norm(ws[f"{col}{f}"].value)
            if (e == t) if exacto else e.startswith(t):
                if not ultima:
                    return f
                hallada = f
        return hallada

    def exige(texto, col="B", desde=1, exacto=False, ultima=False):
        f = fila_con(texto, col, desde, exacto, ultima)
        if f is None:
            raise SystemExit(
                f"\nNo encuentro la etiqueta «{texto}» en la columna {col} de «{HOJA}».\n"
                "La estructura del modelo cambió y el informe NO se generó.\n"
                "Es preferible fallar aquí que publicar cifras leídas de la celda equivocada."
            )
        return f

    def bloque(desde, cols, campos, parar=None, huecos=0):
        """Lee hacia abajo hasta la primera fila vacía (o hasta `parar`).
        `huecos`: filas en blanco que el bloque tiene por dentro y no lo terminan."""
        out, vacias = [], 0
        for f in range(desde, LIM + 1):
            vals = [c(col, f) for col in cols]
            if all(v is None or v == "" for v in vals):
                vacias += 1
                if vacias > huecos:
                    break
                continue
            if parar and norm(vals[0]).startswith(parar.lower()):
                break
            vacias = 0
            out.append(dict(zip(campos, vals)))
        return out

    def no_vacia(desde, col="B"):
        for f in range(desde, LIM + 1):
            if norm(ws[f"{col}{f}"].value):
                return f
        return desde

    # --- localización de cada bloque ---
    f_mes = exige("mes", "B", exige("análisis consolidado", ultima=True), exacto=True)
    mensual = bloque(f_mes + 1, "BCDEFGHIJ",
                     ["mes", "cfo", "capex", "intl", "fcf", "burn", "prewc", "cambio", "lectura"])

    f_ppto = exige("mes", "B", exige("ingresos presupuesto vs ejecutado"), exacto=True)
    f_pte = exige("concepto", "B", exige("puente de caja"), exacto=True)
    f_chq = exige("diferencia (chequeo", "B", f_pte)
    f_intl = exige("país", "F", exige("flujo de caja internacional", "F"), exacto=True)
    f_trm = exige("trm ", "F", f_intl)
    # Bloque opcional: el modelo de agosto 2026 lo eliminó de la hoja. Si no está, la
    # pestaña simplemente no muestra esa tarjeta; no es motivo para no generar el informe.
    f_marg_sec = fila_con("márgenes consolidados", "G")
    f_marg = fila_con("concepto", "G", f_marg_sec, exacto=True) if f_marg_sec else None
    f_dso = exige("dso objetivo")
    f_otros = exige("% recuperado")
    f_res = exige("indicador", "B", exige("resumen ejecutivo consolidado", ultima=True), exacto=True)
    f_msg = exige("mensaje clave", "B", f_res)
    f_dif = exige("diferencia consumo neto", "B", f_res)
    f_proy = exige("mes", "B", exige("runway proyectado"), exacto=True)
    proyeccion = bloque(f_proy + 1, "BCDEF", ["mes", "ingresos", "flujo", "caja", "estado"])
    f_kpis = no_vacia(f_proy + 1 + len(proyeccion) + 1)
    f_conc = exige("conclusión proyectada", "B", f_kpis)

    def dato(etiqueta, col="C", desde=1):
        return c(col, exige(etiqueta, "B", desde))

    # --- corte: se deduce del modelo, no se escribe a mano ---
    LARGO = {"ene": "Enero", "feb": "Febrero", "mar": "Marzo", "abr": "Abril",
             "may": "Mayo", "jun": "Junio", "jul": "Julio", "ago": "Agosto",
             "sep": "Septiembre", "oct": "Octubre", "nov": "Noviembre", "dic": "Diciembre"}
    reales = [x["mes"] for x in mensual
              if isinstance(x.get("mes"), str) and re.fullmatch(r"[A-Za-z]{3}-\d{2}", x["mes"])]
    if not reales:
        raise SystemExit("No pude deducir el mes de corte del bloque mensual.")
    ult = reales[-1]
    corte = f"{LARGO[ult[:3].lower()]} 20{ult[-2:]}"

    datos = {
        "generado": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "modelo": os.path.basename(MODELO),
        "fuente": huella(MODELO),
        "corte": corte,

        # Resumen ejecutivo: caja, burn y runway (B=indicador, C=valor, D=unidad, E=nota)
        "resumen": bloque(f_res + 1, "BCDE", ["indicador", "valor", "unidad", "nota"],
                          parar="mensaje clave"),
        "mensaje": c("C", f_msg),
        "dif_consumo_fcf": {"valor": c("C", f_dif), "unidad": c("D", f_dif), "nota": c("E", f_dif)},

        # Márgenes consolidados del mes de corte (G..K)
        # lleva una fila en blanco por dentro; termina en "Fuente / alcance"
        "margenes": bloque(f_marg + 1, "GHIJK",
                           ["concepto", "colombia", "otros", "consolidado", "criterio"],
                           parar="fuente", huecos=1) if f_marg else [],

        # Evolución mensual del burn (B..J)
        "mensual": mensual,

        # Presupuesto vs ejecutado (B..I)
        "ppto": bloque(f_ppto + 1, "BCDEFGHI",
                       ["mes", "pptoUSD", "ejecUSD", "gapUSD", "cumpl", "pptoCOP", "ejecCOP", "gapCOP"]),

        # Puente de caja del mes (B..D)
        "puente": bloque(f_pte + 1, "BCD", ["concepto", "valor", "comentario"],
                         parar="diferencia (chequeo"),
        "puente_chequeo": {"dif": c("C", f_chq), "estado": c("D", f_chq)},

        # Flujo internacional del mes (F..J)
        "intl": bloque(f_intl + 1, "FGHIJ",
                       ["pais", "ingresosUSD", "egresosUSD", "netoUSD", "netoCOP"]),
        # nombre histórico de la clave; el valor es la TRM del mes de corte
        "trm_junio": c("J", f_trm),

        # Sensibilidad por días de cobro (B..F)
        "dso_actual": {"ingreso_mes": dato("ingreso mensual"), "cxc": dato("cxc clientes actual"),
                       "dso": dato("dso actual clientes"), "caja": dato("caja actual"),
                       "caja_minima": dato("caja mínima operativa")},
        "dso": bloque(f_dso + 1, "BCDEF", ["dso", "cxc", "liberada", "resultante", "estado"]),

        # Palanca: recuperación de otros deudores (B..E)
        "otros_deudores_actual": dato("otros deudores actual"),
        "otros_deudores": bloque(f_otros + 1, "BCDE", ["pct", "liberada", "resultante", "estado"]),

        # Proyección hasta diciembre con el presupuesto (B..F)
        "proyeccion": proyeccion,
        "proyeccion_kpis": bloque(f_kpis, "BCDE", ["indicador", "valor", "unidad", "nota"],
                                  parar="conclusión"),
        "conclusion": c("C", f_conc),
    }

    with open(SALIDA, "w", encoding="utf-8") as f:
        json.dump(datos, f, ensure_ascii=False, indent=1)

    print(f"Informe Burn & Runway extraído del modelo · corte {corte}:")
    for k in ("resumen", "mensual", "ppto", "puente", "intl", "dso", "proyeccion", "proyeccion_kpis", "margenes"):
        print(f"   {k:20s} {len(datos[k])} fila(s)")
    print(f"\nGenerado: {os.path.abspath(SALIDA)}")
    print("Abre el dashboard y ve a la pestaña «Burn & Runway».")


if __name__ == "__main__":
    main()
