# -*- coding: utf-8 -*-
"""Convierte plantillas/Plantilla_Burn_Runway.xlsx en burn_runway.json.

PARA QUE
--------
Hoy burn_runway.json lo produce extraer_burn_runway.py leyendo Modelo_COCO.xlsx
entero. Eso obliga a tener el modelo completo en la carpeta y a que su estructura
no se mueva. Con la plantilla, quien cierra el mes diligencia siete hojas y esto
arma el mismo JSON.

Los dos caminos conviven: si el modelo esta a mano, extraer_burn_runway.py sigue
siendo mas rapido. Este es para cuando no lo esta.

USO
---
    python modelo/plantilla_a_burn_runway.py <plantilla.xlsx>              (en seco)
    python modelo/plantilla_a_burn_runway.py <plantilla.xlsx> --escribir   (aplica)

En seco imprime lo que quedaria y NO toca nada. Con --escribir respalda el JSON
anterior y lo reemplaza.

QUE COMPRUEBA ANTES DE ESCRIBIR
-------------------------------
El tablero busca sus tarjetas por el TEXTO del indicador. Si alguien reescribe
"Caja total consolidada", la tarjeta se queda vacia y nada avisa. Por eso aqui se
exige que las etiquetas que el tablero busca sigan estando, y si falta una, no se
escribe: es preferible fallar aqui que publicar una pestana con huecos.
"""
import json
import os
import shutil
import sys
from datetime import datetime

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)
SALIDA = os.path.join(RAIZ, "burn_runway.json")

# Lo que renderBurnRunway() busca dentro de d.resumen y d.proyeccion_kpis. Si una
# de estas no aparece, la tarjeta correspondiente sale vacia.
EXIGE_RESUMEN = ["caja total", "fcf burn", "consumo neto", "runway consolidado"]
EXIGE_KPIS = ["primer déficit"]


def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, str):
        v = v.strip().replace(".", "").replace(",", ".") if v.count(",") == 1 and v.count(".") > 1 else v.strip()
        try:
            return float(v)
        except ValueError:
            return v          # "N/A" y demas textos del modelo se pasan tal cual
    return float(v)


def _txt(v):
    return None if v is None else str(v).strip() or None


def _filas(ws, primera, cols, clave=2):
    """Lee hacia abajo desde `primera` hasta la primera fila sin clave."""
    out = []
    r = primera
    while True:
        k = ws.cell(r, clave).value
        if k is None or str(k).strip() == "":
            break
        fila = {}
        for nombre, (col, tipo) in cols.items():
            v = ws.cell(r, col).value
            fila[nombre] = _txt(v) if tipo == "t" else _num(v)
        out.append(fila)
        r += 1
    return out


def construir(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    faltan = [h for h in ("CABECERA", "RESUMEN", "MENSUAL", "PRESUPUESTO",
                          "PROYECCION", "PROYECCION_KPIS") if h not in wb.sheetnames]
    if faltan:
        raise SystemExit("\nA la plantilla le faltan hojas: %s\n"
                         "No se genera nada." % ", ".join(faltan))

    cab = wb["CABECERA"]
    corte = _txt(cab["C5"].value)
    modelo = _txt(cab["C6"].value) or os.path.basename(ruta)
    trm = _num(cab["C7"].value)
    mensaje = _txt(cab["C10"].value)
    conclusion = _txt(cab["C11"].value)

    resumen = _filas(wb["RESUMEN"], 6, {
        "indicador": (2, "t"), "valor": (3, "n"), "unidad": (4, "t"), "nota": (5, "t")})
    mensual = _filas(wb["MENSUAL"], 6, {
        "mes": (2, "t"), "cfo": (3, "n"), "capex": (4, "n"), "fcf": (5, "n"),
        "burn": (6, "n"), "cambio": (7, "n"), "lectura": (8, "t")})
    for m in mensual:
        m["intl"] = None
        m["prewc"] = 0
        if m.get("burn") is None:
            m["burn"] = 0
    ppto = _filas(wb["PRESUPUESTO"], 6, {
        "mes": (2, "t"), "pptoCOP": (3, "n"), "ejecCOP": (4, "n"), "gapCOP": (5, "n"),
        "cumpl": (6, "n"), "pptoUSD": (7, "n"), "ejecUSD": (8, "n"), "gapUSD": (9, "n")})
    proyeccion = _filas(wb["PROYECCION"], 6, {
        "mes": (2, "t"), "ingresos": (3, "n"), "flujo": (4, "n"), "caja": (5, "n"),
        "estado": (6, "t")})
    kpis = _filas(wb["PROYECCION_KPIS"], 6, {
        "indicador": (2, "t"), "valor": (3, "n"), "unidad": (4, "t"), "nota": (5, "t")})
    intl = []
    if "INTERNACIONAL" in wb.sheetnames:
        intl = [f for f in _filas(wb["INTERNACIONAL"], 6, {
            "pais": (2, "t"), "ingresosUSD": (3, "n"), "egresosUSD": (4, "n"),
            "netoUSD": (5, "n"), "netoCOP": (6, "n")})
            if f.get("ingresosUSD") is not None or f.get("egresosUSD") is not None]

    # --- controles -------------------------------------------------------
    problemas = []
    etiquetas = " · ".join((f["indicador"] or "").lower() for f in resumen)
    for e in EXIGE_RESUMEN:
        if e not in etiquetas:
            problemas.append("RESUMEN no trae ningun indicador que contenga «%s»: "
                             "esa tarjeta saldria vacia." % e)
    etq = " · ".join((f["indicador"] or "").lower() for f in kpis)
    for e in EXIGE_KPIS:
        if e not in etq:
            problemas.append("PROYECCION_KPIS no trae «%s»: la tarjeta de runway "
                             "hasta el deficit saldria vacia." % e)
    if not any("runway promedio" in (f["indicador"] or "").lower() for f in resumen):
        problemas.append("RESUMEN no trae ningun «Runway promedio ...»: son dos tarjetas.")
    if not corte:
        problemas.append("CABECERA!C5 (Corte) esta vacio: el encabezado diria «—».")
    if not trm:
        problemas.append("CABECERA!C7 (TRM del informe) esta vacia: el tablero no "
                         "puede advertir la diferencia de tasa con el resto.")
    if not mensual:
        problemas.append("MENSUAL no trae meses: el grafico de burn quedaria vacio.")
    if not intl:
        problemas.append("INTERNACIONAL esta vacia: el titulo diria «Burn & Runway · "
                         "Colombia» en vez de «Consolidado (4 paises)».")

    d = {
        "generado": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "modelo": modelo,
        "fuente": {"archivo": os.path.basename(ruta),
                   "bytes": os.path.getsize(ruta),
                   "modificado": datetime.fromtimestamp(
                       os.stat(ruta).st_mtime).strftime("%Y-%m-%d %H:%M"),
                   "origen": "plantilla"},
        "corte": corte,
        "trm_junio": trm,
        "resumen": resumen,
        "mensaje": mensaje,
        "conclusion": conclusion,
        "margenes": [],
        "mensual": mensual,
        "ppto": ppto,
        "puente": [],
        "intl": intl,
        "proyeccion": proyeccion,
        "proyeccion_kpis": kpis,
    }
    return d, problemas


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    escribir = "--escribir" in sys.argv
    if not args:
        print(__doc__)
        return 2
    ruta = args[0]
    if not os.path.exists(ruta):
        print("No encuentro %s" % ruta)
        return 2

    d, problemas = construir(ruta)

    print("\n  Corte: %s · TRM del informe: %s" % (d["corte"], d["trm_junio"]))
    print("  %d indicadores de resumen · %d meses · %d filas de presupuesto · "
          "%d meses proyectados · %d KPIs · %d filiales"
          % (len(d["resumen"]), len(d["mensual"]), len(d["ppto"]),
             len(d["proyeccion"]), len(d["proyeccion_kpis"]), len(d["intl"])))

    if problemas:
        print("\n  CONTROLES EN ALTO — no se escribe nada:")
        for p in problemas:
            print("    · %s" % p)
        return 1

    if not escribir:
        print("\n  Todo en orden. Aplica con:  python %s %s --escribir"
              % (os.path.relpath(__file__, RAIZ).replace("\\", "/"), os.path.basename(ruta)))
        return 0

    if os.path.exists(SALIDA):
        resp = SALIDA.replace(".json", "_antes_%s.json"
                              % datetime.now().strftime("%Y%m%d_%H%M%S"))
        shutil.copy2(SALIDA, resp)
        print("\n  Respaldo: %s" % os.path.basename(resp))
    with open(SALIDA, "w", encoding="utf-8") as f:
        json.dump(d, f, ensure_ascii=False, indent=2)
    print("  Escrito: burn_runway.json")
    print("  Recarga el tablero con Ctrl+F5 para verlo.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
