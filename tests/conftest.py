# -*- coding: utf-8 -*-
"""
Lectura de la base para los tests. Una sola vez por corrida, no una por prueba:
BD_MAESTRA_COCO.xlsx pesa 260 KB y abrirlo 40 veces hace que nadie corra los tests.

La regla que ordena estos dos archivos:

  test_invariantes.py   reglas que NO dependen del mes. Siguen valiendo en agosto,
                        en septiembre y el ano que viene. Si una falla, hay un error
                        de estructura.

  test_cifras.py        las cifras concretas del ultimo cierre. Cambian cada mes, y
                        actualizarlas ES la revision del cierre: obliga a mirar cada
                        numero que se movio y decir por que.
"""
import os
import sys

import pytest

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
# COCO_BASE permite correr la suite contra otra copia de la base. Se usa para la
# prueba de mutacion --romper un valor a proposito y comprobar que los tests lo
# atrapan-- sin tocar la base de verdad.
BASE = os.environ.get("COCO_BASE") or os.path.join(RAIZ, "migracion", "BD_MAESTRA_COCO.xlsx")

CONSOL_SEG = "Consolidación USD"
H1_SEG = "Acumulado H1 2026"
PAISES = ["Colombia", "EE.UU.", "Perú", "Costa Rica"]


@pytest.fixture(scope="session")
def libro():
    import openpyxl
    if not os.path.exists(BASE):
        # Un skip silencioso es peor que un fallo: la suite dice '30 skipped' en verde
        # y nadie lee que no comprobo nada. Si la base la pidio alguien a proposito
        # --COCO_BASE-- y no esta, es un error, no una excusa para no correr.
        if os.environ.get("COCO_BASE"):
            raise AssertionError("COCO_BASE apunta a un archivo que no existe: %s" % BASE)
        pytest.skip("no encuentro %s" % BASE)
    return openpyxl.load_workbook(BASE, data_only=True)


@pytest.fixture(scope="session")
def filas(libro):
    """BD_Indicadores como lista de diccionarios. La cabecera esta en la fila 3."""
    hojas = list(libro["BD_Indicadores"].iter_rows(values_only=True))
    cab = [str(v or "").strip() for v in hojas[2]]
    out = []
    for f in hojas[3:]:
        if not f or not f[0]:
            continue
        out.append({c: f[i] for i, c in enumerate(cab) if c and i < len(f)})
    return out


@pytest.fixture(scope="session")
def trm(libro):
    d = {}
    for f in list(libro["TRM"].iter_rows(values_only=True))[3:]:
        if f and f[0]:
            d[str(f[0])] = f[1]
    return d


@pytest.fixture(scope="session")
def V(filas):
    """Suma de un codigo con los filtros dados. None si no hay ninguna fila.

    Por defecto exige segmento='Consolidación USD' y escenario='Real', que es donde
    vive el P&G oficial. Pasar seg=None desactiva el filtro de segmento -- util para
    comprobar justamente que sin ese filtro la cifra se infla.
    """
    def _v(codigo, periodo=None, pais=None, seg=CONSOL_SEG, escenario="Real", detalle=None):
        t, hubo = 0.0, False
        for f in filas:
            if str(f.get("codigo_indicador") or "").strip() != codigo:
                continue
            if escenario is not None and str(f.get("escenario") or "") != escenario:
                continue
            if periodo is not None and str(f.get("periodo")) != periodo:
                continue
            if pais is not None and str(f.get("pais") or "") != pais:
                continue
            if seg is not None and str(f.get("segmento") or "") != seg:
                continue
            if detalle is not None and str(f.get("detalle") or "") != detalle:
                continue
            v = f.get("valor")
            if v is None:
                continue
            t += float(v)
            hubo = True
        return t if hubo else None
    return _v


@pytest.fixture(scope="session")
def meses_pyg(filas):
    """Los meses que tienen P&G oficial, ordenados. Se deduce de los datos para que
    los tests no haya que tocarlos cuando entre un mes nuevo."""
    m = {str(f.get("periodo")) for f in filas
         if str(f.get("codigo_indicador") or "").strip() == "pyg_ingresos_operacionales"
         and str(f.get("segmento") or "") == CONSOL_SEG
         and str(f.get("escenario") or "") == "Real"}
    return sorted(m)


@pytest.fixture(scope="session")
def ultimo_mes(meses_pyg):
    return meses_pyg[-1] if meses_pyg else None
